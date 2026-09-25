"""Acceptance checks for the CJAP handover artifacts."""
import os, re, sys
from pptx import Presentation
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, ".."))
BULLET = "▪"

print("=" * 74)
print("DECK — CJAP_Handover.pptx")
print("=" * 74)
prs = Presentation(os.path.join(OUT, "CJAP_Handover.pptx"))
over, thin, echo = [], [], []
for i, s in enumerate(prs.slides, 1):
    title, body, has_tbl, has_mono = None, 0, False, False
    body_txt = []
    for sh in s.shapes:
        if sh.has_table:
            has_tbl = True
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if not t:
            continue
        fonts = {r.font.name for p in sh.text_frame.paragraphs for r in p.runs}
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs
                 if r.font.size]
        if sizes and max(sizes) >= 27 and title is None:
            title = t.split("\n")[0]
            continue
        if "Consolas" in fonts:
            has_mono = True
            continue
        if ("Supervaise Inc." in t and "handover" in t) or t.isdigit():
            continue
        if t.startswith("Wireframe from source"):
            continue                                   # figure caption
        clean = t.replace(BULLET, " ")
        body += len(clean.split())
        body_txt.append(clean)
    n = len(s.notes_slide.notes_text_frame.text.split()) if s.has_notes_slide else 0
    flags = []
    if body > 40:
        flags.append("BODY>40"); over.append(i)
    if n < 40:
        flags.append("THIN NOTES"); thin.append(i)
    joined = " ".join(body_txt).lower()
    if title and not has_tbl and not has_mono and joined and \
            title.lower().rstrip(".") in joined and len(joined.split()) < 12:
        flags.append("ECHOES TITLE"); echo.append(i)
    vis = "T" if has_tbl else ("D" if has_mono else "-")
    print(f"{i:>3} body={body:>3} notes={n:>3} [{vis}]  "
          f"{(title or '(title slide)')[:44]:<44} {' '.join(flags)}")
print(f"\nslides={len(prs.slides._sldIdLst)}  over-40={over or 'none'}  "
      f"thin-notes={thin or 'none'}  title-echo={echo or 'none'}")

print()
print("=" * 74)
print("DOCX — accuracy markers")
print("=" * 74)
for f in sorted(x for x in os.listdir(OUT) if x.endswith(".docx")):
    d = Document(os.path.join(OUT, f))
    txt = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                txt += "\n" + c.text
    words = len(txt.split())
    print(f"{f:<42} words={words:>5}  TO CONFIRM={txt.count('[TO CONFIRM')():>2}"
          if False else
          f"{f:<42} words={words:>5}  "
          f"TO-CONFIRM={txt.count('[TO CONFIRM'):>2}  "
          f"Measured={txt.count('Measured'):>2}  Target={txt.count('Target'):>2}  "
          f"tables={len(d.tables):>2}  "
          f"Sources={'Y' if 'Sources:' in txt else 'n'}"
          f" Status={'Y' if 'Status:' in txt else 'n'}")

print()
print("=" * 74)
print("FORBIDDEN-CONTENT CHECK  (Piper outside footnote · Fish Audio · bad corpus size)")
print("=" * 74)
bad = 0
for f in sorted(os.listdir(OUT)):
    path = os.path.join(OUT, f)
    if not os.path.isfile(path):
        continue
    tbl_txt, mono_txt, txt = "", "", ""
    if f.endswith(".docx"):
        d = Document(path)
        txt = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            joined = "\n".join(c.text for r in t.rows for c in r.cells)
            txt += "\n" + joined
            # brand.code_block renders a diagram as a single-cell MONOSPACE table;
            # a single-cell Calibri table is a blockquote callout, not a diagram.
            fonts = {r.font.name for row in t.rows for c in row.cells
                     for p in c.paragraphs for r in p.runs}
            if len(t.rows) == 1 and len(t.columns) == 1 and "Consolas" in fonts:
                mono_txt += "\n" + joined
            elif len(t.rows) > 1:
                tbl_txt += "\n" + joined
    elif f.endswith(".pptx"):
        pr = Presentation(path)
        for sl in pr.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    body = sh.text_frame.text
                    txt += "\n" + body
                    fonts = {r.font.name for pa in sh.text_frame.paragraphs
                             for r in pa.runs}
                    if "Consolas" in fonts:
                        mono_txt += "\n" + body
                if sh.has_table:
                    cells = "\n".join(c.text for r in sh.table.rows for c in r.cells)
                    txt += "\n" + cells
                    tbl_txt += "\n" + cells
            if sl.has_notes_slide:
                txt += "\n" + sl.notes_slide.notes_text_frame.text
    else:
        continue
    issues = []
    # Fish Audio may only appear as an explicit ABSENCE / discrepancy statement.
    for m in re.finditer(r"fish\s*audio", txt, re.I):
        ctx = txt[max(0, m.start() - 200):m.start() + 240].lower()
        if not any(k in ctx for k in ("not ", "no occurrence", "does not exist",
                                      "absent", "never")):
            issues.append("Fish Audio asserted as in use")
            break
    # Piper may only appear as prose/footnote, never inside a table or diagram.
    if "Piper" in tbl_txt:
        issues.append("Piper inside a table")
    if "Piper" in mono_txt:
        issues.append("Piper inside a diagram")
    if re.search(r"1,?109", txt) and "pilot" not in txt.lower():
        issues.append("1,109 corpus cited without pilot framing")
    bad += len(issues)
    print(f"{f:<42} {'OK' if not issues else ' | '.join(issues)}")
print("\nissues:", bad)
