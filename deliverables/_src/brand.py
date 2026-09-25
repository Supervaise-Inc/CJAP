"""Supervaise brand helpers for the CJAP handover artifacts.

Brand identity (colors, banner, company name, tagline) comes from the
supervaise-branding skill. That skill's *structural* rules are written for
single-page certificates and are NOT applied to these reports/decks; only the
identity layer is reused.
"""
import os
from docx.shared import Pt, RGBColor, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SKILL = (r"C:\Users\ASUS\AppData\Roaming\Claude\local-agent-mode-sessions"
         r"\skills-plugin\64b383c4-9d57-417a-9dff-45932f7a509c"
         r"\5a8f8da4-4a7d-4e46-b445-59e548320daf\skills\supervaise-branding")
BANNER = os.path.join(SKILL, "assets", "header_banner.png")

COMPANY = "Supervaise Inc."
TAGLINE = "Unlocking Possibilities Together"

GREEN = RGBColor(0x00, 0xCC, 0x00)
DARK  = RGBColor(0x2D, 0x2D, 0x2D)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY  = RGBColor(0x80, 0x80, 0x80)
GREEN_HEX, DARK_HEX, GREY_HEX = "00CC00", "2D2D2D", "808080"
BAND_HEX, CODE_HEX = "F2F2F2", "F7F7F7"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


# --------------------------------------------------------------- low level
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k if ":" in k else "w:" + k), v)
    return e


def shade(cell, hex_fill):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", fill=hex_fill))


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = _el("w:tblCellMar")
    for name, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        mar.append(_el("w:" + name, w=str(v), type="dxa"))
    tblPr.append(mar)


def borders(table, sz=4, color="BFBFBF"):
    b = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el("w:" + edge, val="single", sz=str(sz), space="0", color=color))
    table._tbl.tblPr.append(b)


def keep_with_next(par):
    par.paragraph_format.keep_with_next = True


# --------------------------------------------------------------- page setup
def setup_page(doc, a4=True, margin=0.75, landscape=False):
    for s in doc.sections:
        if a4:
            s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        if landscape:
            s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Inches(margin)
        s.top_margin = s.bottom_margin = Inches(margin)
        s.header_distance = s.footer_distance = Inches(0.35)
    n = doc.styles["Normal"]
    n.font.name = BODY_FONT
    n.font.size = Pt(11)
    n.font.color.rgb = BLACK
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.line_spacing = 1.02
    rpr = n.element.get_or_add_rPr()
    rpr.append(_el("w:rFonts", ascii=BODY_FONT, hAnsi=BODY_FONT, cs=BODY_FONT))


def banner(doc, width_in=None):
    """Full-bleed-ish brand banner at the top of the body flow."""
    if not os.path.exists(BANNER):
        return False
    sec = doc.sections[0]
    w = width_in or (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(BANNER, width=Inches(w))
    return True


def page_number_footer(doc, left_text):
    """'left text' on the left, 'Page X of Y' on the right."""
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.text = ""
        p.paragraph_format.tab_stops.add_tab_stop(
            sec.page_width - sec.left_margin - sec.right_margin)
        r = p.add_run(left_text + "\t")
        r.font.size, r.font.color.rgb, r.font.name = Pt(8), GREY, BODY_FONT
        for chunk, fld in (("Page ", None), (None, "PAGE"), (" of ", None), (None, "NUMPAGES")):
            if fld:
                run = p.add_run()
                run.font.size, run.font.color.rgb, run.font.name = Pt(8), GREY, BODY_FONT
                f1 = _el("w:fldChar", fldCharType="begin")
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
                it.text = f" {fld} "
                f2 = _el("w:fldChar", fldCharType="end")
                run._r.append(f1); run._r.append(it); run._r.append(f2)
            else:
                run = p.add_run(chunk)
                run.font.size, run.font.color.rgb, run.font.name = Pt(8), GREY, BODY_FONT


def toc_field(doc, levels="1-3"):
    p = doc.add_paragraph()
    run = p.add_run()
    run._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = f' TOC \\o "{levels}" \\h \\z \\u '
    run._r.append(it)
    run._r.append(_el("w:fldChar", fldCharType="separate"))
    t = OxmlElement("w:t")
    t.text = "Right-click and choose “Update Field” to build the table of contents."
    run._r.append(t)
    run._r.append(_el("w:fldChar", fldCharType="end"))
    # make Word rebuild fields on open
    settings = doc.settings.element
    upd = _el("w:updateFields", val="true")
    settings.append(upd)


# --------------------------------------------------------------- text blocks
def heading(doc, text, level=1, size=None, color=None, space_before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = BODY_FONT
    r.font.size = Pt(size or {1: 16, 2: 13, 3: 11.5, 4: 11}[level])
    r.font.color.rgb = color or (DARK if level <= 2 else BLACK)
    if level <= 2:
        pPr = p._p.get_or_add_pPr()
        pb = _el("w:pBdr")
        pb.append(_el("w:bottom", val="single", sz="6", space="2",
                      color=GREEN_HEX if level == 1 else "D9D9D9"))
        pPr.append(pb)
    # outline level so the TOC field can pick it up
    p._p.get_or_add_pPr().append(_el("w:outlineLvl", val=str(level - 1)))
    return p


INLINE_STYLES = ("**", "`")


def rich_runs(par, text, base_size=11, mono_size=None):
    """Render **bold**, `code` and «swap» markers into runs."""
    import re
    mono_size = mono_size or base_size - 1
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|«[^»]+»)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
            r.font.size, r.font.name = Pt(base_size), BODY_FONT
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.size, r.font.name = Pt(mono_size), MONO_FONT
            r.font.color.rgb = DARK
        elif tok.startswith("«") and tok.endswith("»"):
            r = par.add_run(tok); r.bold = True
            r.font.size, r.font.name = Pt(base_size), BODY_FONT
            r.font.color.rgb = RGBColor(0x00, 0x77, 0x00)
        else:
            r = par.add_run(tok)
            r.font.size, r.font.name = Pt(base_size), BODY_FONT
    return par


def para(doc, text, size=11, space_after=4, italic=False, color=None, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    rich_runs(p, text, base_size=size)
    for r in p.runs:
        if italic:
            r.italic = True
        if color is not None:
            r.font.color.rgb = color
    return p


def bullet(doc, text, size=11, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22 + 0.2 * level)
    rich_runs(p, text, base_size=size)
    return p


def numbered(doc, text, size=11):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.26)
    rich_runs(p, text, base_size=size)
    return p


def code_block(doc, text, size=8.5, title=None):
    """Bordered monospace block — survives copy-paste out of Word."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = t.cell(0, 0)
    shade(cell, CODE_HEX)
    borders(t, sz=6, color=DARK_HEX)
    cell_margins(t, 70, 70, 100, 100)
    cell.text = ""
    first = True
    for line in text.split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name, r.font.size, r.font.color.rgb = MONO_FONT, Pt(size), BLACK
        rpr = r._r.get_or_add_rPr()
        rpr.append(_el("w:rFonts", ascii=MONO_FONT, hAnsi=MONO_FONT, cs=MONO_FONT))
    if title:
        cap = doc.add_paragraph()
        cap.paragraph_format.space_before = Pt(1)
        cr = cap.add_run(title)
        cr.font.size, cr.italic, cr.font.color.rgb = Pt(8), True, GREY
    return t


def callout(doc, text, size=10, fill="FFF4CC", edge="BF9000"):
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    shade(cell, fill)
    borders(t, sz=8, color=edge)
    cell_margins(t, 70, 70, 110, 110)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    rich_runs(p, text, base_size=size)
    return t


def table(doc, rows, header=True, size=9, widths=None, first_col_bold=False):
    """rows = list of list of cell strings (markdown inline supported)."""
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    cell_margins(t)
    t.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            txt = row[j] if j < len(row) else ""
            rich_runs(p, txt, base_size=size, mono_size=size - 0.5)
            if header and i == 0:
                shade(cell, DARK_HEX)
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    r.font.size = Pt(size)
            else:
                if (i % 2 == 0) and header:
                    shade(cell, BAND_HEX)
                if first_col_bold and j == 0:
                    for r in p.runs:
                        r.bold = True
        if header and i == 0:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(_el("w:tblHeader", val="true"))
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Inches(w)
    return t


def sources_status(doc, sources, status, size=8.5):
    """The mandatory trailing Sources / Status lines."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    pPr = p._p.get_or_add_pPr()
    pb = _el("w:pBdr")
    pb.append(_el("w:top", val="single", sz="6", space="3", color=GREEN_HEX))
    pPr.append(pb)
    r = p.add_run("Sources: "); r.bold = True; r.font.size = Pt(size)
    r2 = p.add_run(sources); r2.font.size = Pt(size); r2.font.color.rgb = DARK
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r3 = p2.add_run("Status: "); r3.bold = True; r3.font.size = Pt(size)
    r4 = p2.add_run(status); r4.font.size = Pt(size); r4.font.color.rgb = DARK
    return p, p2


def onepager_head(doc, title, purpose, banner_in=4.2):
    banner(doc, width_in=banner_in)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(17), DARK, BODY_FONT
    q = doc.add_paragraph()
    q.paragraph_format.space_after = Pt(7)
    pPr = q._p.get_or_add_pPr()
    pb = _el("w:pBdr")
    pb.append(_el("w:bottom", val="single", sz="10", space="4", color=GREEN_HEX))
    pPr.append(pb)
    rq = q.add_run(purpose)
    rq.italic, rq.font.size, rq.font.color.rgb = True, Pt(10), DARK
