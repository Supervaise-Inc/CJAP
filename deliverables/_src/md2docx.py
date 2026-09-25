"""Markdown -> branded .docx converter for the CJAP handover deliverables.

Reuses the .md already on disk; adds cover page, TOC field, page numbers,
styled tables and bordered monospace code blocks.
"""
import re, sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brand as B
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK


def _split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _is_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def cover(doc, title, subtitle, meta_rows, classification):
    B.banner(doc, width_in=6.77)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb, r.font.name = True, Pt(30), B.DARK, B.BODY_FONT
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(16)
    r2 = p2.add_run(subtitle)
    r2.font.size, r2.font.color.rgb, r2.italic = Pt(13), B.GREY, True
    doc.add_paragraph()
    t = B.table(doc, meta_rows, header=False, size=10, first_col_bold=True)
    for row in t.rows:
        row.cells[0].width = Inches(1.9)
        row.cells[1].width = Inches(4.8)
    for _ in range(2):
        doc.add_paragraph()
    pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = pc.add_run(classification)
    rc.font.size, rc.font.color.rgb, rc.bold = Pt(9), B.GREY, True
    pf = doc.add_paragraph(); pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.paragraph_format.space_after = Pt(0)
    rf = pf.add_run(f"{B.COMPANY} — {B.TAGLINE}")
    rf.font.size, rf.font.color.rgb = Pt(9), B.GREY
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def convert(md_path, out_path, title, subtitle, meta_rows, classification,
            footer_text, drop_h1=True, toc=True):
    text = open(md_path, encoding="utf-8").read()
    doc = Document()
    B.setup_page(doc, a4=True, margin=0.8)
    cover(doc, title, subtitle, meta_rows, classification)

    if toc:
        B.heading(doc, "Contents", level=2, space_before=0)
        B.toc_field(doc)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    lines = text.split("\n")
    i, n = 0, len(lines)
    pending_tbl = []

    def flush_table():
        nonlocal pending_tbl
        if pending_tbl:
            B.table(doc, pending_tbl, header=True, size=8.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            pending_tbl = []

    while i < n:
        line = lines[i]
        s = line.strip()

        # fenced code block -----------------------------------------------
        if s.startswith("```"):
            flush_table()
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            B.code_block(doc, "\n".join(buf).rstrip(),
                         size=7.6 if max((len(x) for x in buf), default=0) > 74 else 8.4)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # tables -----------------------------------------------------------
        if s.startswith("|"):
            if _is_sep(s):
                i += 1; continue
            pending_tbl.append(_split_row(s))
            i += 1; continue
        flush_table()

        # headings ----------------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and drop_h1:
                i += 1; continue
            B.heading(doc, re.sub(r"[*`]", "", txt), level=min(lvl, 4) if lvl > 1 else 1)
            i += 1; continue

        # horizontal rule ----------------------------------------------------
        if re.fullmatch(r"-{3,}", s):
            i += 1; continue

        # blockquote (callout) ------------------------------------------------
        if s.startswith(">"):
            buf = []
            while i < n and (lines[i].strip().startswith(">") or
                             (buf and lines[i].strip() == "")):
                q = lines[i].strip()
                if q.startswith(">"):
                    buf.append(q.lstrip(">").strip())
                elif buf and (i + 1 < n and lines[i + 1].strip().startswith(">")):
                    buf.append("")
                else:
                    break
                i += 1
            body = "\n".join(b for b in buf)
            tbl = doc.add_table(rows=1, cols=1)
            cell = tbl.cell(0, 0)
            B.shade(cell, "EFF7EF"); B.borders(tbl, sz=8, color=B.GREEN_HEX)
            B.cell_margins(tbl, 80, 80, 120, 120)
            cell.text = ""
            first = True
            for bl in body.split("\n"):
                if not bl.strip() and first:
                    continue
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.paragraph_format.space_after = Pt(3)
                B.rich_runs(p, bl.lstrip("- ").strip() if bl.strip().startswith("-") else bl,
                            base_size=9.5)
                if bl.strip().startswith("-"):
                    p.paragraph_format.left_indent = Inches(0.18)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # bullets / numbers ----------------------------------------------------
        mb = re.match(r"^[-*]\s+(.*)$", s)
        mn = re.match(r"^(\d+)\.\s+(.*)$", s)
        if mb:
            B.bullet(doc, mb.group(1), size=10); i += 1; continue
        if mn:
            B.numbered(doc, mn.group(2), size=10); i += 1; continue

        # blank / prose ----------------------------------------------------------
        if not s:
            i += 1; continue
        if s.startswith("*") and s.endswith("*") and len(s) < 60:
            B.para(doc, s.strip("*"), size=10, italic=True, color=B.GREY,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
            i += 1; continue
        B.para(doc, s, size=10.5)
        i += 1

    flush_table()
    B.page_number_footer(doc, footer_text)
    doc.save(out_path)
    return out_path
