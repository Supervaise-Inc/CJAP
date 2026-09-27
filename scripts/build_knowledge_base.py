"""
Phase 6 Step 3 - build knowledge-base/, the folder handed to the Foundation.

knowledge-base/ is a BUILD OUTPUT. It is never hand-edited: corpus/ and data/text/ stay canonical, and this
script derives the folder from them so it cannot drift from the pinned sources.

    python scripts/build_knowledge_base.py            # delete knowledge-base/ and rebuild it
    python scripts/build_knowledge_base.py --check    # rebuild into staging, compare with knowledge-base/, exit 1 on any difference

What it does, in order
  1. Refuses to run unless the four pinned workbooks still match corpus_snapshot.json (verify_pin), and unless the
     corpus, data/text and the pin agree on exactly the same 1,290 document ids (1,295 pinned rows minus the 5
     retired ids in data/csv/retired_doc_ids.csv).
  2. data/text/<id>.md  - bytes copied flat; three declared changes only (see TRANSFORMS below).
  3. data/enriched/*.csv - exported from the pinned workbooks (data sheet only), UTF-8 with BOM, LF, quoted where a
     field holds a comma, sorted by doc id. Every row is re-hashed against the pin before it is written.
  4. corpus/<format>/<theme>/<id>.{md,json} - copied; every card must carry topic_paths with a route_source.
  5. corpus/topic_map.json, voice_card.md, router_prompt.md - copied, LF.
  6. README.md and MANIFEST.md are generated from the data above (counts are computed, never typed).
  7. Final scans: no retired id, no internal batch tag, no CR, no dotfile / __pycache__ / .tmp, no empty directory.

TRANSFORMS (the only differences between knowledge-base/ and the repo sources; each is counted in MANIFEST.md)
  LF        CRLF -> LF on every text file that carries it (153 data/text sources, voice_card.md, router_prompt.md).
  TAG       an internal processing tag at the end of the `Source:` line of 168 book-chapter sources is removed.
  RETIRED   a prose cross-reference to a retired id inside a curated note (one cell, one card) is rewritten to the
            id that superseded it in the retired-documents register. Any OTHER retired-id reference is a hard error.
The corpus content, the workbooks and the dense/sparse indexes are read, never written.
"""
from __future__ import annotations

import argparse
import csv
import datetime
import hashlib
import io
import json
import re
import shutil
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import config  # noqa: E402
from scripts import verify_pin  # noqa: E402
from scripts.build_corpus_snapshot import (  # noqa: E402
    CURATED_COLUMNS, DATA_CSV, SNAPSHOT_PATH, SOURCE_FILES, canon_code, row_hash,
)

OUT_DIR = PROJECT_ROOT / "knowledge-base"
STAGE_DIR = PROJECT_ROOT / "knowledge-base.staging"
CORPUS = PROJECT_ROOT / "corpus"
DATA_TEXT = PROJECT_ROOT / "data" / "text"
INDEX_DIR = PROJECT_ROOT / "data" / "index"
REGISTER = DATA_CSV / "retired_doc_ids.csv"
TAGS_REPORT = PROJECT_ROOT / "reports" / "topic_tags_full_2026-09-27.json"
ROUTING_REPORT = PROJECT_ROOT / "reports" / "routing_check_2026-09-27.json"

# folder, id letter, workbook, csv name, expected rows after retirement (the Phase 6 specification)
FORMATS = [
    ("columns", "C", SOURCE_FILES[0], "columns.csv", 803),
    ("books", "B", SOURCE_FILES[1], "books.csv", 299),
    ("speeches", "S", SOURCE_FILES[2], "speeches.csv", 153),
    ("biography", "G", SOURCE_FILES[3], "biography.csv", 35),
]
FORMAT_NOUN = {"columns": "column", "books": "book chapter", "speeches": "speech", "biography": "biography chapter"}
ROUTE_SOURCES = ("matcher", "affinity", "orphan")
LETTER_FOLDER = {letter: folder for folder, letter, *_ in FORMATS}

# --- declared transforms ---------------------------------------------------------------------------------------------
TAG_RE = re.compile(rb" \(batch-04 split\)")            # data/text only
BATCH_RE = re.compile(rb"batch[-_ ]?0?4", re.I)          # must not exist anywhere in the output
TEXT_SUFFIXES = {".md", ".json", ".csv"}

# Publication years of the books, from the "Book Number Map" sheet of the books workbook. Asserted against the corpus so a
# new work forces an edit here rather than a silent gap in the README.
WORK_YEAR = {
    "Love God, Serve Man": 1994, "Justice and Faith": 1997, "Battles in the Supreme Court": 1998,
    "Leadership by Example: The Davide Standard": 1999, "Transparency, Unanimity & Diversity": 2000,
    "A Centenary of Justice": 2001, "Reforming the Judiciary": 2002, "The Bio-Age Dawns on the Judiciary": 2003,
    "Leveling the Playing Field": 2004, "Judicial Renaissance": 2005, "Liberty and Prosperity": 2006,
    "With Due Respect (Vol. 1)": 2011, "With Due Respect (Vol. 2)": 2011, "With Due Respect (Vol. 3)": 2011,
    "With Due Respect (Vol. 4)": 2011, "With Due Respect (Vol. 5)": 2011, "With Due Respect (Vol. 6)": 2011,
    "With Due Respect (Vol. 7)": 2011,
}
# The 35 biography chapters are two works (BATCH-04_BIOGRAPHY_PROMPTS.md): a range check keeps this honest.
BIOGRAPHY_SPLIT = ("GC001", "GC020", "GC021", "GC035")

# Frozen 40-query retrieval set, measured in Phase 5 (Phase 6 specification, Step 5).
FROZEN_SET_CHANGED, FROZEN_SET_TOTAL, FROZEN_HIT10_DROP_PTS = 14, 40, 2.9

INDEX_PURPOSE = [
    # (file relative to repo root, meta file holding build_date or None, purpose)
    ("data/index/corpus_dense.npy", "data/index/corpus_dense_meta.json",
     "Meaning vectors: one 768-number vector for each of the 13,549 passages (model BAAI/bge-base-en-v1.5). Drives search by meaning."),
    ("data/index/pilot_dense.npy", "data/index/pilot_dense_meta.json",
     "{pilot_dense_note}"),
    ("data/index/pilot_sparse.pkl", None,
     "Keyword index (BM25 over the curated phrases and the passage text). Drives exact-phrase search."),
    ("data/index/corpus_mean.npy", None,
     "The average passage vector. It is subtracted from every vector before two are compared; all three routing thresholds assume this 'centred' scale."),
    ("data/index/topic_centroids.npy", "data/index/topic_centroids_meta.json",
     "Thirty topic vectors, one per dimension: the centred average of that dimension's passages. Used to place a question or a document in a dimension."),
]
JSON_INDEX_PURPOSE = [
    ("data/index/sparse_phrase_dict.json", "Phrase dictionary behind the keyword index."),
    ("data/index/date_index.json", "One entry per document: its date and how precise it is (day, month or year). Books and biography chapters dated 1 January are labelled year."),
    ("corpus/index/chunk_index.json", "Per-document lookup into the passages, with the pin the passages were cut from."),
    ("corpus/index/chunks.jsonl", "The 13,549 passages themselves (every document cut into 200-400 token pieces)."),
]


class BuildError(SystemExit):
    def __init__(self, msg: str):
        super().__init__(f"[build_knowledge_base] REFUSING: {msg}")


def die(msg: str):
    raise BuildError(msg)


# --- small helpers -----------------------------------------------------------------------------------------------------
def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for blk in iter(lambda: f.read(1 << 20), b""):
            h.update(blk)
    return h.hexdigest()


def to_lf(b: bytes) -> tuple[bytes, bool]:
    if b"\r" not in b:
        return b, False
    out = b.replace(b"\r\n", b"\n")
    if b"\r" in out:
        die("a lone carriage return remains after CRLF normalisation - not a line-ending difference, stopping")
    return out, True


def fmt_int(n: int) -> str:
    return f"{n:,}"


def kinds(counter: Counter) -> str:
    return ", ".join(f"{v} {FORMAT_NOUN[k]}{'s' if v != 1 else ''}" for k, v in counter.items() if v)


def load_json(p: Path):
    return json.loads(p.read_text(encoding="utf-8"))


def nice_date(iso: str) -> str:
    d = datetime.date.fromisoformat(iso)
    return f"{d.day} {d.strftime('%B %Y')}"


# --- the retired-documents register ------------------------------------------------------------------------------------
def load_register() -> list[dict]:
    with REGISTER.open("r", encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    if len(rows) != 5:
        die(f"the retired register holds {len(rows)} rows, the Phase 6 specification names 5")
    return rows


def retired_matchers(register: list[dict]):
    ids = [r["doc_id"] for r in register]
    id_re = re.compile(r"(?<![A-Za-z0-9])(" + "|".join(map(re.escape, ids)) + r")(?![0-9])", re.I)
    successor = {r["doc_id"].upper(): r["superseded_by"].strip() for r in register if r["superseded_by"].strip()}
    return ids, id_re, successor


class Rewriter:
    """RETIRED transform: rewrite a reference to a superseded retired id; anything else is left for the final scan to reject."""

    def __init__(self, id_re, successor):
        self.id_re, self.successor = id_re, successor
        self.hits: list[tuple[str, str]] = []           # (where, retired id)

    def text(self, s: str, where: str) -> str:
        def rep(m):
            new = self.successor.get(m.group(1).upper())
            if not new:
                return m.group(0)
            self.hits.append((where, m.group(1)))
            return new
        return self.id_re.sub(rep, s)

    def data(self, b: bytes, where: str) -> bytes:
        return self.text(b.decode("utf-8"), where).encode("utf-8")


# --- staging area ------------------------------------------------------------------------------------------------------
class Stage:
    def __init__(self, root: Path):
        self.root = root
        self.entries: dict[str, tuple[int, str]] = {}

    def emit(self, rel: str, data: bytes):
        if rel in self.entries:
            die(f"duplicate output path {rel}")
        p = self.root / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data)
        self.entries[rel] = (len(data), sha256_bytes(data))


# --- step 1: the gate --------------------------------------------------------------------------------------------------
def gate(register: list[dict]) -> dict:
    if verify_pin.main() != 0:
        die("verify_pin failed - the workbooks no longer match corpus_snapshot.json")
    pinned = load_json(SNAPSHOT_PATH)
    pin_ids = set(pinned["docs"])
    ret_ids = {r["doc_id"] for r in register}
    if not ret_ids <= pin_ids:
        die(f"retired ids not in the pin: {sorted(ret_ids - pin_ids)}")
    expected = pin_ids - ret_ids
    corpus_ids = {p.stem for fmt, *_ in FORMATS for p in (CORPUS / fmt).rglob("*.json")}
    text_ids = {p.stem for p in DATA_TEXT.glob("*.md")}
    for label, got in (("corpus/", corpus_ids), ("data/text/", text_ids)):
        if got != expected:
            die(f"{label} holds {len(got)} ids but pinned rows minus retired is {len(expected)}: "
                f"missing {sorted(expected - got)[:5]}, extra {sorted(got - expected)[:5]}")
    if corpus_ids & ret_ids or text_ids & ret_ids:
        die("a retired id is present in corpus/ or data/text/")
    return {"pin_sha256": sha256_file(SNAPSHOT_PATH), "pinned_docs": pinned["docs"], "pin_ids": pin_ids,
            "expected_ids": expected, "source_files": pinned["source_files"]}


# --- step 3: enriched CSVs ---------------------------------------------------------------------------------------------
def read_workbook(fname: str) -> tuple[list[str], list[tuple[str, tuple]]]:
    wb = openpyxl.load_workbook(DATA_CSV / fname, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]                            # the data sheet - never 'Book Number Map'
    if ws.title == "Book Number Map":
        die("the first sheet of the books workbook is the Book Number Map, expected the data sheet")
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(it) if h not in (None, "")]
    idx = {name: i for i, name in enumerate(header)}
    rows = []
    for row in it:
        ai = idx["Article Code"]
        if ai >= len(row) or row[ai] in (None, ""):
            continue
        rows.append((canon_code(row[ai]), row))
    wb.close()
    return header, rows


def export_csv(fname: str, folder: str, pinned_docs: dict, retired: set, rewriter: Rewriter) -> tuple[bytes, int, list, dict]:
    header, rows = read_workbook(fname)
    want = [c for c in CURATED_COLUMNS if not (folder == "biography" and c == "Link")]
    if header != want:
        die(f"{fname} header is {header}, expected {want}")
    idx = {name: i for i, name in enumerate(header)}
    out_rows = []
    for code, row in rows:
        if code in retired:
            continue
        cells = {c: ("" if idx[c] >= len(row) or row[idx[c]] is None else str(row[idx[c]])) for c in header}
        full = {c: cells.get(c, "") for c in CURATED_COLUMNS}
        if row_hash(full) != pinned_docs[code]:
            die(f"{fname} row {code} does not hash to the pin")
        for c, v in cells.items():
            if "\r" in v or "\n" in v:
                die(f"{fname} {code}.{c} holds a line break; the CSV contract has none")
            cells[c] = rewriter.text(v, f"{folder}.csv {code}.{c}")
        out_rows.append((code, [cells[c] for c in header]))
    out_rows.sort(key=lambda x: x[0])
    ids = [c for c, _ in out_rows]
    if len(ids) != len(set(ids)):
        die(f"{fname} has duplicate doc ids")
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
    w.writerow(header)
    for _, r in out_rows:
        w.writerow(r)
    return ("﻿" + buf.getvalue()).encode("utf-8"), len(out_rows), header, {c: r for c, r in out_rows}


# --- step 4: corpus cards ----------------------------------------------------------------------------------------------
def check_card(d: dict, doc_id: str, letter: str, theme_folder: str, wb_row: dict):
    if d.get("id") != doc_id or d.get("format") != letter or d.get("theme") != theme_folder[0]:
        die(f"{doc_id}: card id/format/theme ({d.get('id')}/{d.get('format')}/{d.get('theme')}) disagrees with its path")
    tp = d.get("topic_paths")
    if not isinstance(tp, dict) or tp.get("route_source") not in ROUTE_SOURCES:
        die(f"{doc_id}: topic_paths carries no valid route_source")
    if tp["route_source"] == "orphan":
        if tp.get("primary") is not None or tp.get("secondary") != []:
            die(f"{doc_id}: an orphan must be primary null / secondary []")
    elif not (isinstance(tp.get("primary"), list) and tp["primary"]):
        die(f"{doc_id}: route_source {tp['route_source']} with an empty primary")
    for card_key, col in (("title", "Title"), ("date", "Date"), ("one_paragraph_summary", "one_paragraph_summary")):
        if d.get(card_key) != wb_row[col]:
            die(f"{doc_id}: card {card_key} differs from the pinned workbook row")
    if (d.get("link") or "") != wb_row.get("Link", ""):
        die(f"{doc_id}: card link differs from the pinned workbook row")


# --- facts ---------------------------------------------------------------------------------------------------------------
def build_facts(cards: dict, topic_map: dict, register: list[dict]) -> dict:
    F: dict = {"n_docs": len(cards)}
    F["by_format"] = Counter(c["_folder"] for c in cards.values())
    F["theme"] = {}
    for c in cards.values():
        t = F["theme"].setdefault(c["theme"], {"label": c["theme_label"], "n": 0})
        if t["label"] != c["theme_label"]:
            die(f"theme {c['theme']} carries two labels")
        t["n"] += 1
    F["route"] = Counter(c["topic_paths"]["route_source"] for c in cards.values())
    F["orphans"] = sorted(i for i, c in cards.items() if c["topic_paths"]["route_source"] == "orphan")
    F["orphans_by_format"] = Counter(cards[i]["_folder"] for i in F["orphans"])
    tags = load_json(TAGS_REPORT)["per_doc"]
    F["below_floor_matcher"] = sorted(i for i, c in cards.items()
                                      if c["topic_paths"]["route_source"] == "matcher" and not tags[i].get("primary"))
    F["bio_below_floor"] = sorted(i for i in cards if cards[i]["_folder"] == "biography"
                                  and (cards[i]["topic_paths"]["route_source"] == "orphan" or i in F["below_floor_matcher"]))
    F["bio_orphans"] = [i for i in F["bio_below_floor"] if cards[i]["topic_paths"]["route_source"] == "orphan"]
    F["bio_matcher"] = [i for i in F["bio_below_floor"] if i not in F["bio_orphans"]]
    if len(F["bio_orphans"]) != 2 or len(F["bio_matcher"]) != 1:
        die("the README wording covers two orphaned and one below-floor biography chapter; the corpus no longer matches")
    F["bio_matcher_primary"] = cards[F["bio_matcher"][0]]["topic_paths"]["primary"][0]
    # books by work
    works: dict[str, list] = defaultdict(list)
    for i, c in cards.items():
        if c["_folder"] == "books":
            works[c["title"].split(" -- ")[0]].append(i)
    if set(works) != set(WORK_YEAR):
        die(f"works in the corpus differ from WORK_YEAR: {sorted(set(works) ^ set(WORK_YEAR))}")
    F["works"] = sorted(((w, WORK_YEAR[w], len(ids)) for w, ids in works.items()), key=lambda x: (x[1], x[0]))
    # biography
    b0, b1, c0, c1 = BIOGRAPHY_SPLIT
    bio = sorted(i for i in cards if cards[i]["_folder"] == "biography")
    if bio[0] != b0 or bio[-1] != c1 or len([i for i in bio if b0 <= i <= b1]) != 20 or len([i for i in bio if c0 <= i <= c1]) != 15:
        die("the biography ids no longer split 20 + 15 as recorded")
    F["bio_titles"] = {i: cards[i]["title"] for i in bio}
    # dates
    for folder in ("columns", "speeches", "biography"):
        ds = sorted(c["date"] for c in cards.values() if c["_folder"] == folder and c["date"])
        F[f"range_{folder}"] = (ds[0], ds[-1])
    F["wdr_chapters"] = sum(n for w, _, n in F["works"] if w.startswith("With Due Respect"))
    jan1 = sorted(i for i, c in cards.items() if c["_folder"] in ("books", "biography") and c["date"].endswith("-01-01"))
    F["jan1"] = jan1
    F["jan1_by_work"] = Counter(("Biography" if cards[i]["_folder"] == "biography" else cards[i]["title"].split(" -- ")[0]) for i in jan1)
    # topic map
    F["topics"] = [(t["id"], t["display_name"], t["definition"], t["doc_count"]) for t in topic_map["topics"].values()]
    F["n_topics"] = len(F["topics"])
    F["intent"] = topic_map["intents"]["robot_identity_meta"]
    members = {d for t in topic_map["topics"].values() for d in t["doc_ids"]}
    if members != {i for i, c in cards.items() if c["topic_paths"]["route_source"] == "matcher"}:
        die("topic_map membership no longer equals the matcher-routed documents")
    F["display"] = {t["id"]: t["display_name"] for t in topic_map["topics"].values()}
    F["retired"] = register
    return F


def routing_paragraph() -> str:
    r = load_json(ROUTING_REPORT)
    sf = r["short_form"]
    n = sf["n"]
    hit1 = sum(1 for x in sf["rows"] if x["hit1"])
    hit3 = sum(1 for x in sf["rows"] if x["hit3"])
    fails = [x for x in sf["rows"] if not x["hit1"]]
    if sorted(x["dimension"] for x in fails) != sorted(sf["never_win_top1"]):
        die("routing report: the never-win list does not match the failing rows")

    if r["never_win_top1"] or r["never_in_top3"]:
        die("the 60-question set now has a dimension that never wins its own question; the README wording says none does - rewrite it")
    oos = config.OUT_OF_SCOPE_THRESHOLD

    def pct(x: float) -> str:
        return (f"{x * 100:.1f}").rstrip("0").rstrip(".") + "%"

    def one(x):
        s = f'`{x["dimension"]}` (“{x["question"]}” went to `{x["top1"]}`'
        if not x["in_scope"]:
            s += f', with a closeness of {x["top1_cos"]}, under the router’s {oos} out-of-scope line'
        return s + ")"
    long_n, long_1, long_3 = r["n_questions"], r["top1_accuracy"], r["top3_accuracy"]
    long_wrong = sum(r["wrong_winners"].values())
    chance = r["chance"]
    words = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}
    twin = next((x for x in fails if x["dimension"] == "twin_beacons_doctrine"), None)
    flag = (" One of the four is the project’s flagship concept, the twin-beacons doctrine.") if twin else ""
    return (
        "**Known limitation – topic routing.** Topic routing is newly rebuilt and has not been validated end to end. "
        "The routing measured here is the closest-dimension router in the app’s retrieval code, which compares a question "
        "with the centre of each of the 30 dimensions. On 30 short questions, one per dimension, it named the intended "
        f"dimension first for {hit1} of {n} ({pct(hit1 / n)}) and within its top three for {hit3} of {n} ({pct(hit3 / n)}). "
        f"{words[len(fails)]} dimensions never won their own question: " + "; ".join(one(x) for x in fails) + f".{flag} "
        f"A second set of {long_n} longer, descriptive questions (two per dimension) did better – {pct(long_1)} first and "
        f"{pct(long_3)} within three, with every dimension winning at least one of its own questions – but {long_wrong} "
        f"questions still went to a wrong first choice. Chance for 30 dimensions is {pct(chance['top1'])} first and "
        f"{pct(chance['top3'])} within three, so the router is doing real work; it is not yet reliable enough to trust unattended. "
        "Two things were not measured: the live robot’s own router (a language-model call driven by router_prompt.md, "
        "which needs an API key that was not available) and end-to-end answer quality. On the frozen 40-query retrieval set, "
        f"{FROZEN_SET_CHANGED} of {FROZEN_SET_TOTAL} selected result sets changed after the move to the 30-dimension map and gold hit@10 "
        f"fell by {FROZEN_HIT10_DROP_PTS} points, while hit@1, hit@3 and hit@5 did not move. Routing was deliberately not adjusted in this delivery. "
        "CE-11 to CE-14 are the gate before the robot is demonstrated."
    )


# --- README --------------------------------------------------------------------------------------------------------------
def render_readme(F: dict, ctx: dict, limitation: str) -> str:
    bf = F["by_format"]
    n = F["n_docs"]
    r = F["route"]
    ob = F["orphans_by_format"]
    works_rows = "\n".join(f"| {w} | {y} | {c} |" for w, y, c in F["works"])
    theme_rows = "\n".join(f"| {k} | {v['label']} | {fmt_int(v['n'])} |" for k, v in sorted(F["theme"].items()))
    topic_rows = "\n".join(f"| `{tid}` | {name.replace('|', chr(92) + '|')} | {defn.replace('|', chr(92) + '|')} | {cnt} |"
                           for tid, name, defn, cnt in F["topics"])
    topic_ids = ", ".join(f"`{t[0]}`" for t in F["topics"])
    ret_rows = "\n".join(f"| {x['title']} | {x['reason']} |" for x in F["retired"])
    bio_lines = "\n".join(f"- `{i}` – {t}" for i, t in F["bio_titles"].items() if i in F["bio_below_floor"])
    o1, o2 = F["bio_orphans"]
    (m1,) = F["bio_matcher"]
    c_lo, c_hi = F["range_columns"]
    s_lo, s_hi = F["range_speeches"]
    jan_book = sum(v for k, v in F["jan1_by_work"].items() if k != "Biography")
    jan_bio = F["jan1_by_work"].get("Biography", 0)
    jan_book_works = sorted((k, v) for k, v in F["jan1_by_work"].items() if k != "Biography")
    jan_works = ("all from *" + jan_book_works[0][0] + "*") if len(jan_book_works) == 1 else \
        "from " + ", ".join(f"*{k}* ({v})" for k, v in jan_book_works)
    return f"""# The Panganiban Knowledge Base

This folder is the knowledge base behind the conversation app that speaks as retired Philippine Chief Justice
Artemio V. Panganiban (“CJP”). It was prepared by the Supervaise project team for the Foundation. It holds
{fmt_int(n)} documents – his newspaper columns, the chapters of his books, his speeches and two biographies – each in
three forms, plus the map the app uses to decide which documents to read when someone asks a question.

Everything here is a copy. It is produced by a script from the project’s pinned source files and is never edited by hand; if
something in this folder is wrong, the fix is made at the source and the folder is rebuilt. `MANIFEST.md` lists every file with its
SHA-256 fingerprint, so any copy can be checked against the original.

## 1. What the corpus is and where it came from

| Kind of document | Documents | What they are |
|---|---:|---|
| Columns | {fmt_int(bf['columns'])} | *With Due Respect*, Justice Panganiban’s column in the Philippine Daily Inquirer, {nice_date(c_lo)} to {nice_date(c_hi)} |
| Book chapters | {fmt_int(bf['books'])} | One document per chapter, from {len(F['works'])} works (table below) |
| Speeches | {fmt_int(bf['speeches'])} | Speeches and addresses, {nice_date(s_lo)} to {nice_date(s_hi)} |
| Biography chapters | {fmt_int(bf['biography'])} | Chapters of two biographies written about him by other authors |
| **Total** | **{fmt_int(n)}** | |

**Book chapters by work.** “Chapters” are the chapters (or sections) that were turned into documents.

| Work | Published | Chapters in the corpus |
|---|---:|---:|
{works_rows}

The seven *With Due Respect* volumes reprint columns from 2007 to 2011, so they are the only column material from those years in the
corpus; the columns of that period that were never sourced are described in section 6. *Liberty and Prosperity* (2006) is his own book; it is not the biography with a
similar title described next.

**Biography chapters.** `GC001` to `GC020` are the twenty chapters, Prologue to Epilogue, of Reginald T. Yu’s *Liberty and
Prosperity: The Making of Chief Justice Artemio Villaseñor Panganiban Jr.*; `GC021` to `GC035` are the fifteen chapters of a second
biography (Chapter 1, Early Years, to Chapter 15, Papal Award). In these documents the author is the biographer, not Justice Panganiban,
and his own words appear only as quotations.

**How each document was prepared.** The source text of every document was collected and cleaned first. A description of each was then
written in a fixed set of fields – a one-paragraph summary, keywords, the people and institutions named, his stances, anecdotes, the
audience it suits, the register (tone) it is written in, and the decision-making signals it shows – with AI assistance and spot-checks by the project team, and
recorded in four spreadsheets. Those descriptions are the project’s notes *about* a document; they are not Justice Panganiban’s words.
The words themselves are in `data/text`.

## 2. What is in this folder

```
knowledge-base/
  README.md, MANIFEST.md
  data/
    text/         {fmt_int(n)} source texts, one file per document, named <doc_id>.md
    enriched/     columns.csv, books.csv, speeches.csv, biography.csv - the descriptive fields, one row per document
  corpus/
    columns/ books/ speeches/ biography/   one folder per theme; each document is a pair: <doc_id>.md and <doc_id>.json
    topic_map.json     the 30-dimension topic map (section 5)
    voice_card.md      how the app is told to speak as Justice Panganiban
    router_prompt.md   how the app is told to choose a dimension for a question
```

* **`data/text/<doc_id>.md`** – the source text. Columns, book chapters and speeches start with a short header (title, date, publisher, source); biography chapters are the chapter text alone.
* **`data/enriched/*.csv`** – the four spreadsheets exported as plain CSV files (UTF-8, with a byte-order mark so Excel opens them correctly). The column called *Article Code* is the document’s ID. The biography file has no *Link* column because chapters of a book have no web address.
* **`corpus/<kind>/<theme>/<doc_id>.md`** – the document’s *card*: its descriptive fields at the top, then the full text.
* **`corpus/<kind>/<theme>/<doc_id>.json`** – the same descriptive fields in a form a program can read, plus `topic_paths` (section 5) and `source_xlsx`, which names the spreadsheet the row came from (the spreadsheets themselves are not in this folder; `data/enriched` replaces them).

The search indexes (the files that let the app look documents up by meaning and by keyword) are large binary files and are **not** in this
folder. They stay in the project repository under `data/index/`; `MANIFEST.md` names each one, its size and its build date.

## 3. What a document ID means

An ID such as `SB085` has three parts:

1. **First letter – the kind:** `C` column, `B` book chapter, `S` speech, `G` biography chapter.
2. **Second letter – the theme:** `A` to `E` (section 4).
3. **Three digits – a number** within that kind and theme. The number is only a label; it does not follow date order.

**IDs are never reused.** If a document is withdrawn, its ID is retired for good and no later document takes it. That is why the
numbers within a kind and theme have gaps, and why five IDs are missing (section 6).

## 4. The five themes

Each document sits in exactly one theme, which is its folder and its second letter. Themes are broad shelves; they are not the topic map.

| Letter | Theme | Documents |
|---|---|---:|
{theme_rows}

## 5. The topic map

The **topic map** (`corpus/topic_map.json`) is the list of subjects the app can route a question to. Each subject is a **dimension** – for
example *Libel and Cybercrime* or *Liberty and Prosperity: the Twin Beacons*. When someone asks a question, the app picks the one or two dimensions that fit it and reads
the documents filed under them.

The current map has **{F['n_topics']} dimensions**. It was derived from the whole {fmt_int(n)}-document corpus in September 2026. It replaces a 35-topic
list that had been written when the corpus held only 79 documents and had never seen a book; that list is retired and is not in this folder.

Every dimension has a definition, a list of member documents and a set of keywords that identify it. Beside the {F['n_topics']} dimensions the map holds one **intent**,
`robot_identity_meta`, for questions about the robot itself rather than about the law; it is handled separately and is not a subject.

| Dimension ID | Name | What it covers | Documents listed in the map |
|---|---|---|---:|
{topic_rows}

**How each document is filed.** Every card carries a `topic_paths` field with a `primary` dimension, up to a few `secondary` ones, and a `route_source` that says how the filing was decided:

| `route_source` | Meaning | Documents |
|---|---|---:|
| `matcher` | Keyword matchers (the dimension’s own keywords, names and cases) recognised the document. | {fmt_int(r['matcher'])} |
| `affinity` | No matcher recognised it, but its passages sit close to one dimension. *Affinity* is the similarity between a document’s closest passage and the centre of a dimension; a document is filed there only above a fixed floor (0.31 on the app’s centred scale). | {fmt_int(r['affinity'])} |
| `orphan` | Neither. The `primary` is an explicit `null` – no topic – never a blank. | {fmt_int(r['orphan'])} |
| **Total** | | **{fmt_int(sum(r.values()))}** |

`topic_map.json` lists the {fmt_int(r['matcher'])} documents the matchers filed. The {fmt_int(r['affinity'])} affinity-filed documents carry their dimension in their own `.json` card only, so
adding up the member lists in the map gives fewer documents than the corpus holds.

## 6. What is deliberately not here

* **About 222 Inquirer columns from February 2007 to April 2011.** They were never sourced, so there is nothing to put in the corpus. The columns that are here start on {nice_date(c_lo)}. Some columns from that period do appear, as chapters of the seven *With Due Respect* volumes ({F['wdr_chapters']} chapters); the ~222 unsourced columns are not otherwise represented.
* **Five retired documents.** Each was withdrawn from the corpus in September 2026. Their IDs are retired and are deliberately not listed in this folder.

| Retired document | Why it was withdrawn |
|---|---|
{ret_rows}

## 7. Known limitations

**Three biography chapters do not fit the dimensions.**

{bio_lines}

No dimension came close enough to these chapters to clear the affinity floor. `{o1}` and `{o2}` carry no topic (`primary` is `null`); `{m1}` carries a keyword-matched
filing ({F['display'][F['bio_matcher_primary']]}) although it sits below that floor. Across the whole corpus {r['orphan']} documents have no topic ({ob['columns']} columns, {ob['books']} books, {ob['speeches']} speeches, {ob['biography']} biography chapters; the full list is in `MANIFEST.md`),
and {len(F['below_floor_matcher'])} more are filed by a keyword matcher although their affinity is below the floor.

**Dates: 1 January means “the year”.** {len(F['jan1'])} documents – {jan_bio} biography chapters and {jan_book} book chapters ({jan_works}) – are dated 1 January of a year. For these the project has no date finer than the year, and 1 January is a placeholder, not a day.
The date is left as it appears in the source, everywhere in this folder. The app’s date index labels these documents “year” precision instead of “day”, so a search by date does not treat them as exact days. Please read any 1 January date on a book or biography chapter as “sometime that year”.

**The descriptive fields are drafted, not authored.** Summaries, keywords, stances and the rest were drafted with AI assistance and spot-checked, not read line by line by Justice Panganiban or the Foundation. Treat them as a guide to the text, and the text in `data/text` as the authority.

{limitation}
"""


def render_manifest(F: dict, ctx: dict, limitation: str, entries: dict[str, tuple[int, str]]) -> str:
    bf = F["by_format"]
    n = F["n_docs"]
    counts = ctx["counts"]
    pinned_rows = sum(ctx["pin_rows"].values())
    tot_files = len(entries) + 1
    tot_bytes = sum(v[0] for v in entries.values())
    ret_rows = "\n".join(f"| {x['title']} | {x['type']} | {x['reason']} | {x['superseded_by'] or chr(8212)} | {x['retired_on']} |"
                         for x in F["retired"])
    idx_rows = "\n".join(f"| `{p}` | {fmt_int(sz)} | {date} | {purpose} |" for p, sz, date, purpose in ctx["indexes"])
    jidx_rows = "\n".join(f"| `{p}` | {fmt_int(sz)} | {date} | {purpose} |" for p, sz, date, purpose in ctx["json_indexes"])
    orphan_rows = ", ".join(f"`{i}`" for i in F["orphans"])
    below = ", ".join(f"`{i}`" for i in F["below_floor_matcher"])
    files = "\n".join(f"| `{p}` | {fmt_int(entries[p][0])} | `{entries[p][1]}` |" for p in sorted(entries))
    sf = ctx["source_files"]
    pin_rows = "\n".join(f"| `{f['filename']}` | {f['row_count']} | `{f['sha256']}` |" for f in sf)
    lf = ctx["xform"]
    return f"""# MANIFEST – knowledge-base/

Every file in this folder, its size and its SHA-256, plus the counts and the exact list of ways these files differ from the project
repository. `README.md` explains what the corpus is; this file is the audit trail. `MANIFEST.md` cannot list itself.

Totals: **{fmt_int(tot_files)} files** ({fmt_int(len(entries))} listed below plus this manifest), **{fmt_int(tot_bytes)} bytes** without this manifest.

## Built from

`scripts/build_knowledge_base.py`, which refuses to run unless the four pinned workbooks still match `corpus_snapshot.json`
(SHA-256 `{ctx['pin_sha256']}`), and unless the corpus, the source texts and the pin all name the same {fmt_int(n)} documents.

| Pinned workbook | Rows | SHA-256 |
|---|---:|---|
{pin_rows}

## Counts

| Area | Columns | Books | Speeches | Biography | Total |
|---|---:|---:|---:|---:|---:|
| `data/text/<doc_id>.md` | {counts['text']['columns']} | {counts['text']['books']} | {counts['text']['speeches']} | {counts['text']['biography']} | {fmt_int(sum(counts['text'].values()))} |
| `data/enriched/*.csv` rows | {counts['csv']['columns']} | {counts['csv']['books']} | {counts['csv']['speeches']} | {counts['csv']['biography']} | {fmt_int(sum(counts['csv'].values()))} |
| `corpus/**/<doc_id>.md` cards | {counts['md']['columns']} | {counts['md']['books']} | {counts['md']['speeches']} | {counts['md']['biography']} | {fmt_int(sum(counts['md'].values()))} |
| `corpus/**/<doc_id>.json` cards | {counts['json']['columns']} | {counts['json']['books']} | {counts['json']['speeches']} | {counts['json']['biography']} | {fmt_int(sum(counts['json'].values()))} |

Also: `corpus/topic_map.json`, `corpus/voice_card.md`, `corpus/router_prompt.md`, `README.md` and this file.

**The arithmetic.** The four pinned workbooks hold {fmt_int(pinned_rows)} rows ({" + ".join(str(v) for v in ctx['pin_rows'].values())}). Five documents are retired, so
{fmt_int(pinned_rows)} − 5 = **{fmt_int(n)}** documents are live. `data/text/` holds one file per live document, so it holds {fmt_int(n)} files. The
retired documents have no file in `data/text/` to remove, so the five are not subtracted a second time.

## Differences from the project repository

The delivered files are byte-for-byte the repository files except for the three changes below. Nothing else was altered.

1. **Line endings.** Every file here uses LF. {lf['lf_data_text']} of the {fmt_int(n)} source texts ({kinds(lf['lf_kinds'])}) use CRLF in the repository, where they are exact copies of their sources; in this folder they are LF. `voice_card.md` and `router_prompt.md` were also CRLF in the repository ({lf['lf_voice']} files) and are LF here.
2. **An internal processing tag removed from {lf['tag_files']} source texts ({kinds(lf['tag_kinds'])}).** At the end of the `Source:` header line these files carried a parenthesised note recording which build stage split the chapter out. It is a project-internal note, not part of the source; it is removed so old and new documents are indistinguishable. Nothing else in those files changes.
3. **One cross-reference rewritten.** A curated note in speech `SC090` (field `register_markers`) named the retired duplicate of speech `SB085`. In `data/enriched/speeches.csv` and in `corpus/speeches/C_biographical_personal/SC090.json` it now names `SB085`, the surviving copy, as the retired-documents register records. This is the only descriptive text in the folder that differs from the pinned workbook (two places: the CSV cell and its card).

## Retired documents (excluded everywhere)

The five documents below are not in any folder, spreadsheet or list in this folder, and their IDs do not appear anywhere in it.

| Title | Kind | Reason | Superseded by | Retired on |
|---|---|---|---|---|
{ret_rows}

## Search indexes (not in this folder)

The binary indexes stay in the project repository under `data/index/`. They are large, rebuilt by scripts, and not part of this delivery. Sizes are exact.
The build date is taken from the index’s own metadata file where it records one and from the file’s modified date where it does not.

| File | Bytes | Built | What it is for |
|---|---:|---|---|
{idx_rows}

Companion lookup tables (JSON), also kept in the repository:

| File | Bytes | Built | What it is for |
|---|---:|---|---|
{jidx_rows}

## Topic filing

Route sources across the {fmt_int(n)} cards: matcher {fmt_int(F['route']['matcher'])}, affinity {fmt_int(F['route']['affinity'])}, orphan {fmt_int(F['route']['orphan'])} (sum {fmt_int(sum(F['route'].values()))}).

Documents with no topic (`primary` is `null`): {orphan_rows}.

Documents filed by a keyword matcher although their affinity to any dimension is below the 0.31 floor ({len(F['below_floor_matcher'])}): {below}.

{limitation}

## Files

| Path | Bytes | SHA-256 |
|---|---:|---|
{files}
"""


# --- indexes -----------------------------------------------------------------------------------------------------------
def index_rows() -> tuple[list, list]:
    def built(rel: str, meta_rel: str | None) -> str:
        if meta_rel:
            m = load_json(PROJECT_ROOT / meta_rel)
            if m.get("build_date"):
                return m["build_date"]
        return datetime.datetime.fromtimestamp((PROJECT_ROOT / rel).stat().st_mtime).date().isoformat()

    rows = []
    dense_same = None
    for rel, meta, purpose in INDEX_PURPOSE:
        p = PROJECT_ROOT / rel
        if not p.exists():
            die(f"{rel} is missing - rebuild the indexes before delivering")
        if rel.endswith("pilot_dense.npy"):
            dense_same = sha256_file(p) == sha256_file(PROJECT_ROOT / "data/index/corpus_dense.npy")
            purpose = ("An exact copy of corpus_dense.npy under the file name the running app loads." if dense_same
                       else "Meaning vectors under the file name the running app loads (differs from corpus_dense.npy).")
        rows.append((rel, p.stat().st_size, built(rel, meta), purpose))
    jrows = []
    for rel, purpose in JSON_INDEX_PURPOSE:
        p = PROJECT_ROOT / rel
        if not p.exists():
            die(f"{rel} is missing - rebuild the indexes before delivering")
        jrows.append((rel, p.stat().st_size, built(rel, None), purpose))
    return rows, jrows


# --- the build ---------------------------------------------------------------------------------------------------------
def build(stage_dir: Path) -> Stage:
    register = load_register()
    retired_ids, id_re, successor = retired_matchers(register)
    rewriter = Rewriter(id_re, successor)
    g = gate(register)
    pinned_docs = g["pinned_docs"]
    retired = set(retired_ids)

    if stage_dir.exists():
        shutil.rmtree(stage_dir)
    stage_dir.mkdir()
    st = Stage(stage_dir)
    xform = {"lf_data_text": 0, "lf_voice": 0, "lf_other": 0, "tag_files": 0, "lf_kinds": Counter(), "tag_kinds": Counter()}
    counts = {k: {} for k in ("text", "csv", "md", "json")}

    # -- data/text
    n_text = Counter()
    for doc_id in sorted(g["expected_ids"]):
        raw = (DATA_TEXT / f"{doc_id}.md").read_bytes()
        b, crlf = to_lf(raw)
        xform["lf_data_text"] += crlf
        b2, k = TAG_RE.subn(b"", b)
        if k > 1:
            die(f"{doc_id}: the processing tag occurs {k} times")
        xform["tag_files"] += k
        b2 = rewriter.data(b2, f"data/text/{doc_id}.md")
        st.emit(f"data/text/{doc_id}.md", b2)
        kind = LETTER_FOLDER[doc_id[0]]
        n_text[kind] += 1
        xform["lf_kinds"][kind] += crlf
        xform["tag_kinds"][kind] += k
    counts["text"] = dict(n_text)

    # -- enriched CSVs (rows keyed by folder for the card cross-check)
    wb_rows: dict[str, dict] = {}
    pin_rows: dict[str, int] = {}
    for folder, letter, fname, csv_name, want in FORMATS:
        data, n, header, rows = export_csv(fname, folder, pinned_docs, retired, rewriter)
        if n != want:
            die(f"{csv_name} would hold {n} rows, the specification says {want}")
        st.emit(f"data/enriched/{csv_name}", data)
        counts["csv"][folder] = n
        wb_rows.update({c: dict(zip(header, r)) for c, r in rows.items()})
        pin_rows[folder] = sum(1 for d in pinned_docs if d[0] == letter)

    # -- corpus cards
    cards: dict[str, dict] = {}
    n_md, n_json = Counter(), Counter()
    for folder, letter, *_ in FORMATS:
        for theme_dir in sorted(p for p in (CORPUS / folder).iterdir() if p.is_dir()):
            if folder == "biography" and theme_dir.name != "C_biographical_personal":
                die(f"biography holds a theme folder other than C_biographical_personal: {theme_dir.name}")
            for jp in sorted(theme_dir.glob("*.json")):
                doc_id = jp.stem
                mp = jp.with_suffix(".md")
                if not mp.exists():
                    die(f"{doc_id}: card .json without its .md")
                if doc_id in retired:
                    die(f"{doc_id}: a retired id is in corpus/")
                jb, crlf_j = to_lf(jp.read_bytes())
                mb, crlf_m = to_lf(mp.read_bytes())
                xform["lf_other"] += crlf_j + crlf_m
                jb = rewriter.data(jb, f"corpus/{folder}/{theme_dir.name}/{doc_id}.json")
                mb = rewriter.data(mb, f"corpus/{folder}/{theme_dir.name}/{doc_id}.md")
                d = json.loads(jb.decode("utf-8"))
                check_card(d, doc_id, letter, theme_dir.name, wb_rows[doc_id])
                d["_folder"] = folder
                cards[doc_id] = d
                st.emit(f"corpus/{folder}/{theme_dir.name}/{doc_id}.json", jb)
                st.emit(f"corpus/{folder}/{theme_dir.name}/{doc_id}.md", mb)
                n_json[folder] += 1
                n_md[folder] += 1
            extra = [p.name for p in theme_dir.iterdir() if p.suffix not in (".md", ".json")]
            if extra:
                die(f"{theme_dir}: unexpected files {extra}")
    counts["json"], counts["md"] = dict(n_json), dict(n_md)
    if set(cards) != g["expected_ids"]:
        die("the card set is not the expected id set")

    # -- voice
    for name in ("topic_map.json", "voice_card.md", "router_prompt.md"):
        b, crlf = to_lf((CORPUS / "voice" / name).read_bytes())
        xform["lf_voice"] += crlf
        st.emit(f"corpus/{name}", rewriter.data(b, f"corpus/{name}"))
    topic_map = json.loads((stage_dir / "corpus" / "topic_map.json").read_text(encoding="utf-8"))
    members = {d for t in topic_map["topics"].values() for d in t["doc_ids"]}
    if not members <= set(cards):
        die("topic_map.json names documents that are not in the corpus")

    # -- README + MANIFEST
    F = build_facts(cards, topic_map, register)
    for jr in F["retired"]:
        jr["title"] = jr["title"].strip()
    indexes, json_indexes = index_rows()
    declared = {("speeches.csv SC090.register_markers", "SA085"), ("corpus/speeches/C_biographical_personal/SC090.json", "SA085")}
    if set(rewriter.hits) != declared or len(rewriter.hits) != len(declared):
        die(f"retired-id references found beyond the one declared case (SC090): {rewriter.hits}; the MANIFEST wording covers only that case")
    xform["rewrites"] = len(rewriter.hits)
    ctx = {"counts": counts, "pin_rows": pin_rows, "pin_sha256": g["pin_sha256"], "source_files": g["source_files"],
           "indexes": indexes, "json_indexes": json_indexes, "xform": xform}
    limitation = routing_paragraph()
    st.emit("README.md", render_readme(F, ctx, limitation).encode("utf-8"))
    st.emit("MANIFEST.md", render_manifest(F, ctx, limitation, dict(st.entries)).encode("utf-8"))
    final_scans(st, id_re)
    return st


def final_scans(st: Stage, id_re):
    root = st.root
    on_disk = sorted(p.relative_to(root).as_posix() for p in root.rglob("*") if p.is_file())
    if on_disk != sorted(st.entries):
        die("files on disk differ from files emitted")
    for p in root.rglob("*"):
        if p.is_dir() and not any(p.iterdir()):
            die(f"empty directory {p.relative_to(root)}")
        parts = p.relative_to(root).parts
        for part in parts:
            if part.startswith(".") or part == "__pycache__" or part.endswith((".tmp", ".pyc")):
                die(f"forbidden name {'/'.join(parts)}")
        if BATCH_RE.search(p.name.encode()) or id_re.search(p.name):
            die(f"forbidden token in the file name {'/'.join(parts)}")
    for rel in on_disk:
        b = (root / rel).read_bytes()
        if b"\r" in b:
            die(f"{rel}: carries a carriage return")
        if BATCH_RE.search(b):
            die(f"{rel}: carries the internal batch tag")
        m = id_re.search(b.decode("utf-8"))
        if m:
            die(f"{rel}: a retired id ({m.group(0)}) reaches the knowledge base")


def swap(st: Stage):
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    st.root.rename(OUT_DIR)


def check(st: Stage) -> int:
    if not OUT_DIR.exists():
        print("[build_knowledge_base] --check: knowledge-base/ does not exist")
        return 1
    have = {p.relative_to(OUT_DIR).as_posix(): sha256_file(p) for p in OUT_DIR.rglob("*") if p.is_file()}
    want = {rel: h for rel, (_, h) in st.entries.items()}
    bad = sorted(set(have) ^ set(want)) + sorted(k for k in set(have) & set(want) if have[k] != want[k])
    if bad:
        print(f"[build_knowledge_base] --check: knowledge-base/ has drifted from the sources ({len(bad)} files), e.g. {bad[:5]}")
        return 1
    print(f"[build_knowledge_base] --check: knowledge-base/ matches a fresh build ({len(want)} files)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("--check", action="store_true", help="rebuild into staging and compare with knowledge-base/; change nothing")
    args = ap.parse_args()
    try:
        st = build(STAGE_DIR)
        if args.check:
            rc = check(st)
        else:
            swap(st)
            total = sum(v[0] for v in st.entries.values())
            print(f"[build_knowledge_base] wrote {OUT_DIR.name}/: {len(st.entries)} files, {total:,} bytes")
            rc = 0
    except BaseException:
        if STAGE_DIR.exists():
            shutil.rmtree(STAGE_DIR, ignore_errors=True)
        raise
    if STAGE_DIR.exists():
        shutil.rmtree(STAGE_DIR, ignore_errors=True)
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
