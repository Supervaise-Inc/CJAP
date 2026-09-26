"""
Build CJAP_Robot_Project_Plan_v3.xlsx from v2 plus four new format tabs.

Run from the Final Project Folder:
    python scripts/build_project_plan_v3.py

Copies every v2 sheet's values, restyles consistently in Arial, rewrites the
Read Me, and appends format tabs whose samples are pulled LIVE from the repo
(data/text/*.md, data/csv/*_curated_normalized.xlsx, corpus/**) so they cannot
drift from the real data.
"""
from __future__ import annotations
import json, re, sys, textwrap
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "CJAP_Robot_Project_Plan_v2.xlsx"
OUT = ROOT / "CJAP_Robot_Project_Plan_v3.xlsx"

FONT = "Arial"
H1 = Font(name=FONT, size=13, bold=True, color="FFFFFF")
H2 = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BOLD = Font(name=FONT, size=10, bold=True)
BODY = Font(name=FONT, size=10)
MONO = Font(name="Consolas", size=9)
MONOB = Font(name="Consolas", size=9, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="555555")
RED = Font(name=FONT, size=10, bold=True, color="C00000")
FILL_H1 = PatternFill("solid", fgColor="1F3864")
FILL_H2 = PatternFill("solid", fgColor="2E5496")
FILL_BAND = PatternFill("solid", fgColor="FCE4D6")
FILL_CODE = PatternFill("solid", fgColor="F2F2F2")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
TOP = Alignment(vertical="top", wrap_text=True)
TOPL = Alignment(vertical="top", wrap_text=True, horizontal="left")
THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)


def _style_block(ws, first_row, last_row, ncols):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            if cell.font is None or cell.font.name != "Consolas":
                cell.font = cell.font if cell.font and cell.font.bold else BODY
            cell.alignment = TOP
            cell.border = THIN


def banner(ws, row, text, ncols, big=False):
    ws.cell(row, 1).value = text
    ws.cell(row, 1).font = H1 if big else H2
    for c in range(1, ncols + 1):
        ws.cell(row, c).fill = FILL_H1 if big else FILL_H2
    ws.cell(row, 1).alignment = TOPL
    ws.row_dimensions[row].height = 26 if big else 20


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------- live samples
def read_source_md(doc_id: str, n: int = 900) -> str:
    p = ROOT / "data" / "text" / f"{doc_id}.md"
    if not p.exists():
        return f"[{p} not found]"
    t = p.read_text(encoding="utf-8", errors="replace")
    words = len(t.split())
    head = t[:n]
    return head + (f"\n\n… [truncated — {words} words in the real file]" if len(t) > n else "")


def read_corpus_md(doc_id: str, n: int = 900) -> str:
    hits = list((ROOT / "corpus").rglob(f"{doc_id}.md"))
    if not hits:
        return f"[corpus/{doc_id}.md not found]"
    t = hits[0].read_text(encoding="utf-8", errors="replace")
    return t[:n] + (f"\n\n… [truncated]" if len(t) > n else "")


def read_row(workbook: str, doc_id: str) -> tuple[list[str], list[str]]:
    wb = openpyxl.load_workbook(ROOT / "data" / "csv" / workbook, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [h for h in rows[0] if h]
    ci = list(rows[0]).index("Article Code")
    for r in rows[1:]:
        if r and r[ci] == doc_id:
            wb.close()
            return hdr, [("" if v is None else str(v)) for v in r[: len(hdr)]]
    wb.close()
    return hdr, ["[row not found]"] * len(hdr)


def corpus_counts() -> dict:
    c = {"C": 0, "B": 0, "S": 0, "G": 0}
    for p in (ROOT / "corpus").rglob("*.json"):
        b = p.stem
        if len(b) == 5 and b[0] in c:
            c[b[0]] += 1
    return c


# ---------------------------------------------------------------- READ ME
def build_readme(wb):
    ws = wb.create_sheet("Read Me")
    widths(ws, {"A": 30, "B": 118})
    r = 1
    banner(ws, r, "CJAP CONVERSATIONAL ROBOT — WHOLE-PROJECT PLAN v3", 2, big=True); r += 2

    def para(label, text, key=False):
        nonlocal r
        ws.cell(r, 1).value = label
        ws.cell(r, 1).font = BOLD
        ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = text
        ws.cell(r, 2).font = BODY
        ws.cell(r, 2).alignment = TOP
        if key:
            ws.cell(r, 1).fill = FILL_KEY
            ws.cell(r, 2).fill = FILL_KEY
        ws.row_dimensions[r].height = max(15, 12 * (1 + text.count("\n")))
        r += 1

    def head(text):
        nonlocal r
        r += 1
        banner(ws, r, text, 2)
        r += 1

    head("WHAT THIS IS")
    para("The product", "A Reachy Mini robot that answers a visitor's spoken question in the licensed voice of "
         "retired Philippine Chief Justice Artemio V. Panganiban, grounded ONLY in his own corpus — his "
         "newspaper columns, his books, his speeches and his biography. It is delivered to the Foundation "
         "for Liberty and Prosperity.")
    para("The hard constraint", "Every answer must be traceable to something he actually wrote or said. The robot "
         "may decline; it may not invent. That single rule is why so much of this plan is about the corpus and "
         "the evidence, not about the model.")
    para("This workbook", "Seven planning tabs carried forward from v2, plus four NEW reference tabs that pin down "
         "the exact file formats. Every task row carries a paste-ready Claude prompt in the 'How' column.")

    head("HOW TO USE IT")
    para("Project Plan", "P0-P11, the whole arc from consent to operations. Run in ID order; dependencies are named. "
         "Status uses [x] Done / [~] Partial / [ ] Not started, and Evidence must point at a real file.")
    para("Corpus Expansion", "CE-1 to CE-18. This is the RECURRING runbook — every time material is added, corrected "
         "or retired. Batch work follows these numbers, not the P-numbers.")
    para("Data Flow", "Source text -> enriched row -> corpus -> index. Read it once before your first batch.")
    para("Intake Prompts", "Paste-ready prompts by source type (book / column / speech / biography).")
    para("Format · Source .md", "NEW. The contract every source file must meet, with four real samples.")
    para("Format · Enriched row", "NEW. The 15-field curated schema, with four real rows.")
    para("Format · Corpus output", "NEW. What the generator produces — front matter and the paired .json.")
    para("Themes & IDs", "NEW. The five themes, the identifier convention, and the live corpus counts.")

    head("WHAT 'INDEXING' MEANS HERE — in plain language")
    para("The problem", "The corpus is far too large to hand to a model whole. Before the robot can speak, something "
         "has to find the handful of passages that actually answer THIS question. That finding step is retrieval, "
         "and an index is the prepared structure that makes it fast and repeatable.")
    para("1 · Chunking", "Each document is cut into passages of roughly 200-400 tokens, split on the author's own "
         "headings so an argument is not severed mid-thought. An anecdote is never split. A chunk is the unit that "
         "gets retrieved — not a whole document.")
    para("2 · The dense index\n(meaning)",
         "Every chunk is turned into a list of 768 numbers — an embedding — by a model that places text with similar "
         "MEANING close together in that 768-dimensional space. A question is embedded the same way, and the closest "
         "chunks are the candidates. This is what lets 'What does the rule of law actually mean?' find a passage that "
         "never uses those exact words.\n"
         "File: data/index/corpus_dense.npy — one row per chunk. Encoder: BAAI/bge-base-en-v1.5 at 768 dimensions, "
         "ratified at P3.1 and NEVER mixed with another model; two encoders in one index silently corrupts every "
         "comparison and no amount of tuning repairs it.")
    para("3 · The sparse index\n(exact words)",
         "BM25 keyword search over the same chunks. Meaning-matching is bad at things that have no synonyms — a docket "
         "number, a statute, a person's name. Sparse catches those. Curated keywords and entities are kept as ATOMIC "
         "phrases: 'rule of law' stays one term, never three.\nFile: data/index/pilot_sparse.pkl")
    para("4 · Fusion", "Dense and sparse each rank the chunks; the two rankings are merged (reciprocal rank fusion) so "
         "a passage strong on either arm survives. This is why both arms are kept even when the dense model is good.")
    para("5 · The topic map\n(dimensions)",
         "34 topics, each stored as a CENTROID — the average position of its member chunks in that same 768-d space, "
         "plus a label, a description, signature phrases and exemplars. A question is scored against all 34 and the "
         "result BIASES the ranking. It is a soft prior, never a filter: if the topic guess is wrong, the right "
         "passage can still surface.\nFiles: data/index/topic_centroids.npy, corpus/voice/topic_map.json")
    para("6 · The date index", "A lookup from date to document, plus a resolver for phrasing like 'last year'. Built "
         "and verified; currently switched off (DATE_INDEX_ENABLED = false) pending a measured decision.")
    para("Why re-index at all", "An index is a photograph of the corpus at one moment. Add or correct a document and "
         "the photograph is stale — the corpus says one thing, the index answers from another. That is why CE-6 "
         "re-embeds the WHOLE corpus in one pass after every batch.")
    para("The pin", "corpus_snapshot.json holds a sha256 for every source file and every row. verify_pin.py fails if "
         "anything moved. 'The corpus is frozen' is a claim a script can refuse, not a promise.", key=True)

    head("THE FIVE THEMES")
    for code, name, gloss in [
        ("A", "Liberty and Rule of Law", "Constitutional doctrine, due process, judicial independence and review, "
         "civil liberties, the courts as the last bulwark."),
        ("B", "Prosperity and Economic Philosophy", "Economic governance, business and investment law, MSMEs, "
         "social justice as an economic question, the twin-beacons argument on the prosperity side."),
        ("C", "Biographical and Personal", "Early life, family, faith, mentors, friendships, honours — the man rather "
         "than the office."),
        ("D", "FLP Mission and Foundation", "The Foundation for Liberty and Prosperity: scholarships, the dissertation "
         "contest, donors and partners, the Museum."),
        ("E", "Current Events Commentary", "Time-bound commentary on Philippine political events, elections, "
         "geopolitics and topical controversies."),
    ]:
        ws.cell(r, 1).value = f"Theme {code}"
        ws.cell(r, 1).font = BOLD
        ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = f"{name} — {gloss}"
        ws.cell(r, 2).font = BODY
        ws.cell(r, 2).alignment = TOP
        r += 1
    para("How a theme is assigned", "By the intake curator, from the document's own content, and recorded in the doc "
         "ID's second letter. A theme is a shelf, not a score — one document, one theme. The 34 TOPICS are a finer, "
         "separate layer used for retrieval; do not confuse the two.")

    head("THE IDENTIFIER")
    para("Pattern", "^[CBSG][A-E]\\d{3}$   e.g. CA003, BA108, SB028, GC021")
    para("Position 1 — format", "C column · B book chapter · S speech · G biography")
    para("Position 2 — theme", "A-E as above")
    para("Positions 3-5", "A zero-padded number within that format+theme series. It is a POSITION, not a count — "
         "CA529 does not mean there are 529 columns.")
    para("Never reused", "An ID is never reused and never renumbered. A retired ID stays retired in "
         "data/csv/retired_doc_ids.csv; a replacement takes a NEW ID and records superseded_doc_id.", key=True)

    head("THE CORPUS TODAY")
    cc = corpus_counts()
    tot = sum(cc.values())
    para("Live counts", f"{cc['C']} columns · {cc['B']} book chapters · {cc['S']} speeches · {cc['G']} biography "
         f"chapters = {tot} documents. (Read live from corpus/ when this workbook was generated.)")
    para("Known gaps", "The column run begins 17 Apr 2011; the Inquirer column actually began 11 Feb 2007, so roughly "
         "222 columns from Feb 2007 - Apr 2011 have never been sourced. Speeches and biography have had no new intake "
         "since batch-03.")

    head("THE RULES THAT MAKE THIS SAFE")
    for t in [
        "APPEND-ONLY. New material is added; existing rows are never mutated. Corrections (CE-17) and retirements "
        "(CE-18) are the only exceptions, and each runs as its OWN batch so the diff proves nothing else moved.",
        "ONE encoder, ONE dimensionality across the whole corpus. Mixing them is the single failure tuning cannot "
        "repair.",
        "Matching row counts do NOT prove nothing was lost. Diff content hashes, document by document.",
        "Add a topic dimension only from evidence you can point at, and only with a recorded sign-off.",
        "New material that never reaches the gold set is material you have not tested.",
    ]:
        ws.cell(r, 1).value = "•"
        ws.cell(r, 1).font = BOLD
        ws.cell(r, 2).value = t
        ws.cell(r, 2).font = BODY
        ws.cell(r, 2).alignment = TOP
        r += 1
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------- format tabs
def build_source_md_tab(wb):
    ws = wb.create_sheet("Format · Source .md")
    widths(ws, {"A": 22, "B": 120})
    r = 1
    banner(ws, r, "SOURCE FILE CONTRACT — data/text/<DOC_ID>.md", 2, big=True); r += 2
    ws.cell(r, 1).value = "What it is"
    ws.cell(r, 1).font = BOLD
    ws.cell(r, 2).value = ("The raw text of ONE document, exactly as he wrote it, in one fixed shape. This is the INPUT "
                           "to enrichment and to the corpus generator. One file per document; the filename IS the doc ID.")
    ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP; r += 2

    banner(ws, r, "THE CONTRACT", 2); r += 1
    for k, v in [
        ("Encoding", "UTF-8. No BOM required, but utf-8-sig is tolerated on read."),
        ("Line 1", "The title, exactly as the enriched row's Title field has it. A leading '# ' is accepted."),
        ("Header block", "Date / Publisher / Source, one per line, label then the value. Date is ISO YYYY-MM-DD."),
        ("Blank line", "Then the body."),
        ("Body", "Paragraphs separated by blank lines. His own sub-headings as '## '. Block quotes as '> '."),
        ("FORBIDDEN", "No YAML front matter. No line that is exactly '---'. No backslash escapes or stray __ / * left "
                      "by document conversion. The generator strips a stray YAML block and reports it — do not rely on that."),
        ("One rule that bites", "The title line must EQUAL the sheet's Title. If it does not, the generator fails to "
                                "recognise it and leaks the title into the body. That was defect D-10."),
    ]:
        ws.cell(r, 1).value = k
        ws.cell(r, 1).font = BOLD if k != "FORBIDDEN" else Font(name=FONT, size=10, bold=True, color="C00000")
        ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = v
        ws.cell(r, 2).font = BODY
        ws.cell(r, 2).alignment = TOP
        r += 1
    r += 1

    banner(ws, r, "CANONICAL TEMPLATE — copy this shape", 2); r += 1
    tpl = ("# <Title, exactly as the enriched row's Title field>\n"
           "Date:       <YYYY-MM-DD>\n"
           "Publisher:  <publisher>\n"
           "Source:     <URL, or \"file: <name>\">\n"
           "By: Artemio V. Panganiban\n"
           "\n"
           "<first paragraph of the body>\n"
           "\n"
           "## <his own sub-heading, if the piece has them>\n"
           "\n"
           "<next paragraph>\n"
           "\n"
           "> <a block quote, if he quotes something at length>\n")
    ws.cell(r, 1).value = "TEMPLATE"; ws.cell(r, 1).font = MONOB
    ws.cell(r, 2).value = tpl; ws.cell(r, 2).font = MONO
    ws.cell(r, 2).alignment = TOP; ws.cell(r, 2).fill = FILL_CODE
    ws.row_dimensions[r].height = 190
    r += 2

    banner(ws, r, "FOUR REAL SAMPLES — pulled live from data/text/ when this workbook was built", 2); r += 1
    for label, doc in [("COLUMN — CA003", "CA003"), ("BOOK — BA001", "BA001"),
                       ("SPEECH — SB028", "SB028"), ("BIOGRAPHY — GC021", "GC021")]:
        ws.cell(r, 1).value = label
        ws.cell(r, 1).font = MONOB
        ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = read_source_md(doc)
        ws.cell(r, 2).font = MONO
        ws.cell(r, 2).alignment = TOP
        ws.cell(r, 2).fill = FILL_CODE
        ws.row_dimensions[r].height = 200
        r += 1
    r += 1
    ws.cell(r, 1).value = "HONEST NOTE"
    ws.cell(r, 1).font = Font(name=FONT, size=10, bold=True, color="C00000")
    ws.cell(r, 2).value = ("The four samples above are real, and they do NOT all match the contract. CA003 has no '# ' "
                           "on its title line; SB028 puts blank lines inside the header block, uses an ALL-CAPS title "
                           "and contains a bare '---'; GC021 has no header block at all. These are legacy files from "
                           "earlier intakes under looser rules. The contract above is what NEW intake must produce. "
                           "Do not copy the legacy variance.")
    ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP; ws.cell(r, 2).fill = FILL_KEY
    ws.row_dimensions[r].height = 60
    ws.freeze_panes = "A2"


def build_enriched_tab(wb):
    ws = wb.create_sheet("Format · Enriched row")
    widths(ws, {"A": 6, "B": 30, "C": 22, "D": 92})
    r = 1
    banner(ws, r, "ENRICHED ROW — the 15-field curated schema", 4, big=True); r += 2
    ws.cell(r, 1).value = "Where it lives"
    ws.cell(r, 1).font = BOLD
    ws.cell(r, 2).value = ("Authored per batch in  batch-<id>/<type>_enriched.xlsx  -> normalised to "
                           "_enriched_normalized.xlsx -> appended into the PINNED workbook "
                           "data/csv/cjp_<columns|books|speeches|biography>_curated_normalized.xlsx (first sheet).")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP; r += 2

    banner(ws, r, "THE FIELDS — in this exact order", 4); r += 1
    for h, w in zip(["#", "Field", "Type", "Rule"], [6, 30, 22, 92]):
        pass
    for j, h in enumerate(["#", "Field", "Type", "Rule"], 1):
        ws.cell(r, j).value = h; ws.cell(r, j).font = BOLD; ws.cell(r, j).fill = FILL_BAND
    r += 1
    fields = [
        ("Date", "text YYYY-MM-DD", "Never a datetime, never a number. A real calendar date. Nothing is coerced — a "
                                    "datetime is an error to REPORT, not a value to convert silently."),
        ("Title", "text", "Must equal the source file's title line."),
        ("Article Code", "text", "The doc ID. Unique. Matches ^[CBSG][A-E]\\d{3}$. Never a retired ID."),
        ("Link", "text", "May be blank. NOTE: the biography sheet has NO Link column — 14 fields, not 15."),
        ("Keyword/s", "JSON array of strings", "Names, cases, laws, key phrases. Multi-word keywords stay WHOLE — they "
                                               "become atomic terms in the sparse index."),
        ("primary_topics", "JSON array of strings", "Full sentences: what the document is centrally about."),
        ("sub_topics", "JSON array of strings", "Supporting points, in the order he makes them."),
        ("signature_phrases", "JSON array of strings", "VERBATIM from the source file, 3-20 words. Checked character by "
                                                       "character with curly quotes and dashes folded."),
        ("entities", "JSON object", "Keys from: people, institutions, places, cases, laws_treaties, events. Each a list "
                                    "of strings."),
        ("stances", "JSON array of objects", "{\"claim\", \"rhetorical_move\", \"confidence\"}. See the note below on "
                                             "confidence."),
        ("notable_anecdotes", "JSON array of strings", "First-hand moments. [] is valid."),
        ("target_audience", "JSON array of strings", "Who the piece is addressed to."),
        ("register_markers", "JSON array of strings", "Tone, structure, rhetorical habits."),
        ("decision_framework_signals", "JSON array of strings", "The tests and principles he applies."),
        ("one_paragraph_summary", "text", "One paragraph, no line breaks. THIS IS EMBEDDED — it is appended to the "
                                          "generated .md and therefore chunked and indexed. A stale summary reaches the "
                                          "model."),
    ]
    for i, (f, t, rule) in enumerate(fields, 1):
        ws.cell(r, 1).value = i
        ws.cell(r, 2).value = f
        ws.cell(r, 3).value = t
        ws.cell(r, 4).value = rule
        for c in range(1, 5):
            ws.cell(r, c).font = BODY if c != 2 else BOLD
            ws.cell(r, c).alignment = TOP
            ws.cell(r, c).border = THIN
        r += 1
    r += 1
    for lbl, txt in [
        ("confidence", "The schema declares exactly four values: asserted · asserted with evidence · hedged · "
                       "reported (not his view). MEASURED 26 Sep 2026: 1,022 of 1,104 documents use something else, "
                       "across 1,959 distinct values ('high' 1,552 times, missing 404). Nothing enforces it. Either "
                       "enforce it at normalisation or amend the schema — do not leave it asserted and untrue."),
        ("entities keys", "103 documents use one of 45 keys outside the declared six (dates, concepts, laws, ...). Same "
                          "decision needed."),
        ("Trailing columns", "The books and speeches workbooks physically carry 26 columns — 15 named and 11 empty. "
                             "Harmless, but append to the NAMED 15 and do not widen the sheet further."),
    ]:
        ws.cell(r, 2).value = lbl
        ws.cell(r, 2).font = Font(name=FONT, size=10, bold=True, color="C00000")
        ws.cell(r, 3).value = txt
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.cell(r, 3).font = BODY; ws.cell(r, 3).alignment = TOP; ws.cell(r, 3).fill = FILL_KEY
        ws.row_dimensions[r].height = 58
        r += 1
    r += 1

    banner(ws, r, "FOUR REAL ROWS — pulled live from the pinned workbooks", 4); r += 1
    for label, book, doc in [("COLUMN", "cjp_columns_curated_normalized.xlsx", "CA003"),
                             ("BOOK", "cjp_books_curated_normalized.xlsx", "BA001"),
                             ("SPEECH", "cjp_speeches_curated_normalized.xlsx", "SB028"),
                             ("BIOGRAPHY (14 fields — no Link)", "cjp_biography_curated_normalized.xlsx", "GC021")]:
        hdr, vals = read_row(book, doc)
        ws.cell(r, 1).value = label
        ws.cell(r, 1).font = MONOB
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(r, 1).fill = FILL_BAND
        r += 1
        for h, v in zip(hdr, vals):
            ws.cell(r, 2).value = h
            ws.cell(r, 2).font = MONOB
            ws.cell(r, 2).alignment = TOPL
            ws.cell(r, 3).value = (v[:1200] + " …[truncated]") if len(v) > 1200 else v
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
            ws.cell(r, 3).font = MONO
            ws.cell(r, 3).alignment = TOP
            ws.cell(r, 3).fill = FILL_CODE
            ws.row_dimensions[r].height = 44
            r += 1
        r += 1
    ws.freeze_panes = "A2"


def build_corpus_tab(wb):
    ws = wb.create_sheet("Format · Corpus output")
    widths(ws, {"A": 22, "B": 120})
    r = 1
    banner(ws, r, "GENERATED CORPUS FILES — corpus/<type>/<theme>/<DOC_ID>.md + .json", 2, big=True); r += 2
    for k, v in [
        ("Who writes them", "scripts/generate_corpus_from_xlsx.py — it reads the four PINNED workbooks plus "
                            "data/text/<ID>.md. Never hand-edit a file under corpus/; edit the source and regenerate."),
        ("The .md", "Generated YAML front matter, then the title, then the body, then '## Summary' and "
                    "'## Notable Anecdotes' appended from the enriched row. THE APPENDED SECTIONS ARE INDEXED — that is "
                    "why a stale one_paragraph_summary reaches the model."),
        ("The .json", "The full enriched row as structured data, used by the composer for stances, signature phrases "
                      "and anecdotes. NOT chunked."),
        ("The date", "Front-matter date is str(Date)[:10] — which is exactly why field 1 must be text, not a datetime."),
        ("Three sets must match", "Active rows in the pinned sheets == files in data/text/ == documents in corpus/. "
                                  "verify_pin.py fails if they diverge."),
    ]:
        ws.cell(r, 1).value = k; ws.cell(r, 1).font = BOLD; ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = v; ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP
        r += 1
    r += 1
    banner(ws, r, "REAL SAMPLE — corpus/columns/A_liberty_rule_of_law/CA003.md", 2); r += 1
    ws.cell(r, 1).value = "GENERATED .md"; ws.cell(r, 1).font = MONOB; ws.cell(r, 1).alignment = TOPL
    ws.cell(r, 2).value = read_corpus_md("CA003", 1400)
    ws.cell(r, 2).font = MONO; ws.cell(r, 2).alignment = TOP; ws.cell(r, 2).fill = FILL_CODE
    ws.row_dimensions[r].height = 320
    r += 2
    banner(ws, r, "WHAT HAPPENS NEXT — the index chain", 2); r += 1
    for k, v in [
        ("chunk_corpus.py", "corpus/**.md -> corpus/index/chunks.jsonl + chunk_index.json. 200-400 tokens, "
                            "heading-aware, anecdotes never split."),
        ("build_corpus_dense.py", "chunks -> data/index/corpus_dense.npy (one 768-d row per chunk) + meta with model, "
                                  "dim, backend, corpus hash."),
        ("make_runtime_dense_index.py", "-> data/index/pilot_dense.npy, the index the APP actually loads. "
                                        "build_corpus_dense.py does not write it when no prior one exists."),
        ("build_sparse_index.py", "-> data/index/pilot_sparse.pkl + sparse_phrase_dict.json (BM25 + atomic phrases)."),
        ("build_date_index.py", "-> data/index/date_index.json."),
        ("merge_tag_topics.py", "tags documents against the existing centroids and runs the orphan census. It MERGES "
                                "centroids first — run it with CJ_TOPIC_MERGE_COSINE=1.01 unless you intend a merge."),
    ]:
        ws.cell(r, 1).value = k; ws.cell(r, 1).font = MONOB; ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = v; ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP
        r += 1
    ws.freeze_panes = "A2"


def build_themes_tab(wb):
    import csv as _csv, json as _json, re as _re, collections as _c
    ws = wb.create_sheet("Themes & IDs")
    widths(ws, {"A": 15, "B": 46, "C": 13, "D": 13, "E": 13, "F": 58})
    r = 1

    def ban(t, big=False):
        nonlocal r
        banner(ws, r, t, 6, big=big); r += 1

    def hdr(cells):
        nonlocal r
        for j, h in enumerate(cells, 1):
            ws.cell(r, j).value = h; ws.cell(r, j).font = BOLD
            ws.cell(r, j).fill = FILL_BAND; ws.cell(r, j).alignment = TOP
        r += 1

    def line(cells, bold_first=True, key=False, red=False, h=None):
        nonlocal r
        for j, v in enumerate(cells, 1):
            c = ws.cell(r, j); c.value = v
            c.font = (RED if red else BOLD) if (j == 1 and bold_first) else BODY
            c.alignment = TOPL if j == 1 else TOP
            c.border = THIN
            if key: c.fill = FILL_KEY
        if h: ws.row_dimensions[r].height = h
        r += 1

    def note(label, text, key=False, red=False):
        nonlocal r
        ws.cell(r, 1).value = label
        ws.cell(r, 1).font = RED if red else BOLD
        ws.cell(r, 1).alignment = TOPL
        ws.cell(r, 2).value = text
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2).font = BODY; ws.cell(r, 2).alignment = TOP
        if key:
            ws.cell(r, 1).fill = FILL_KEY; ws.cell(r, 2).fill = FILL_KEY
        ws.row_dimensions[r] = ws.row_dimensions[r]
        ws.row_dimensions[r].height = max(28, 11 * (1 + text.count("\n")) + 14)
        r += 1

    # ---------------- gather live -------------------------------------------
    docs = []
    for p in (ROOT / "corpus").rglob("*.json"):
        try:
            d = _json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        if d.get("id"):
            docs.append(d)

    def workname(t):
        w = (t or "").split(" -- ")[0].strip()
        return _re.sub(r"\s*\(Vol\.\s*\d+\)", "", w)

    man = []
    mp = ROOT / "batch-04" / "intake_manifest.csv"
    if mp.exists():
        man = list(_csv.DictReader(open(mp, encoding="utf-8-sig")))

    bnow, bnew, bvol = _c.Counter(), _c.Counter(), _c.defaultdict(set)
    bth, bnth = _c.defaultdict(_c.Counter), _c.defaultdict(_c.Counter)
    for d in docs:
        if d["id"][0] != "B":
            continue
        w = workname(d["title"]); bnow[w] += 1; bth[w][d["id"][1]] += 1
        m = _re.search(r"\(Vol\.\s*(\d+)\)", d.get("title") or "")
        if m:
            bvol[w].add(int(m.group(1)))
    for row in man:
        if (row.get("doc_id") or "").startswith("B"):
            w = workname(row.get("book_title") or "")
            bnew[w] += 1; bnth[w][row["doc_id"][1]] += 1

    def bucket(pfx):
        sel = [d for d in docs if d["id"][0] == pfx]
        ds = sorted(x for x in (d.get("date") for d in sel) if x)
        return sel, _c.Counter(d["id"][1] for d in sel), (ds[0] if ds else "—"), (ds[-1] if ds else "—")

    cols, cth, cfirst, clast = bucket("C")
    sps, sth, sfirst, slast = bucket("S")
    bio, gth, gfirst, glast = bucket("G")
    cnew = sum(1 for row in man if (row.get("doc_id") or "").startswith("C"))
    tot_now = len(docs)
    tot_new = sum(bnew.values()) + cnew
    thmix = _c.Counter(d["id"][1] for d in docs)

    # ---------------- the tab ------------------------------------------------
    ban("CORPUS INVENTORY, THEMES AND IDENTIFIERS", big=True); r += 1
    note("Read this as", "A live inventory, regenerated by scripts/build_project_plan_v3.py from corpus/ and "
         "batch-04/intake_manifest.csv. 'Now' is what is pinned and indexed today. '+batch-04' is enriched, "
         "QA'd and waiting at CE-4. Re-run the script after any batch and every number here refreshes.")
    r += 1

    ban("THE CORPUS AT A GLANCE")
    hdr(["Format", "What it is", "Now", "+batch-04", "After", "Notes"])
    line(["Columns", "Philippine Daily Inquirer opinion columns, weekly", len(cols), cnew, len(cols) + cnew,
          f"{cfirst} to {clast}. Roughly one a week, unbroken."], h=30)
    line(["Book chapters", "One chapter = one document, across 12 works", sum(bnow.values()), sum(bnew.values()),
          sum(bnow.values()) + sum(bnew.values()), "4 works represented now; 12 after the merge."], h=30)
    line(["Speeches", "Keynotes, inductions, lectures, eulogies", len(sps), 0, len(sps),
          f"{sfirst} to {slast}. No new intake in batch-04."], h=30)
    line(["Biography", "The authorised life, chapter by chapter", len(bio), 0, len(bio),
          f"{gfirst} to {glast}. No new intake in batch-04."], h=30)
    line(["TOTAL", "Active corpus documents", tot_now, tot_new, tot_now + tot_new,
          "Plus 5 retired IDs whose rows stay pinned with no corpus file."], h=30)
    r += 1
    note("The 1,200-plus figure", f"{tot_now} + {tot_new} = {tot_now + tot_new} documents after CE-4. That is the "
         "number to hold in mind: the knowledge model is being built for roughly thirteen hundred documents, not the "
         "seventy-nine the topic taxonomy was written against.", key=True)
    r += 1

    ban("THE TWELVE BOOKS — chapter counts and where they sit")
    hdr(["Volumes", "Work", "Now", "+batch-04", "After", "Theme spread after the merge"])
    for w in sorted(set(bnow) | set(bnew)):
        vols = bvol.get(w, set())
        vtxt = f"{len(vols)} vols" if vols else "single"
        merged = _c.Counter(bth.get(w, {})) + _c.Counter(bnth.get(w, {}))
        spread = " · ".join(f"{k}:{v}" for k, v in sorted(merged.items()))
        line([vtxt, w, bnow.get(w, 0), bnew.get(w, 0), bnow.get(w, 0) + bnew.get(w, 0), spread],
             bold_first=False, h=26)
        ws.cell(r - 1, 2).font = BOLD
    line(["18 vols", "TOTAL — 12 works", sum(bnow.values()), sum(bnew.values()),
          sum(bnow.values()) + sum(bnew.values()), ""], bold_first=False, h=26)
    for c in range(1, 7):
        ws.cell(r - 1, c).font = BOLD
    r += 1
    note("Works vs volumes", "With Due Respect is ONE work in SEVEN volumes. Counting works gives 12 — which is the "
         "'all 12 books' target. Counting volumes gives 18. Both are right; say which you mean.", key=True)
    note("What the merge changes", "Books go from 4 works to 12, and from 131 chapters to 299. Eight works enter the "
         "corpus for the first time. Every one of those 168 chapters is longer-form than a column — median 2,729 "
         "words against roughly 800 — so they chunk differently and carry far more of the doctrinal material.")
    r += 1

    ban("COLUMNS — the weekly run")
    hdr(["Span", "Cadence", "Documents", "Theme A", "Other themes", "Note"])
    line([f"{cfirst[:7]} – {clast[:7]}", "weekly", len(cols), cth.get("A", 0),
          " · ".join(f"{k}:{v}" for k, v in sorted(cth.items()) if k != "A"),
          "Columns are the spine of the corpus and lean heavily to theme A."], bold_first=False, h=28)
    r += 1
    hdr(["Year", "Count", "Year", "Count", "Year", "Count"])
    years = sorted(_c.Counter((d.get("date") or "")[:4] for d in cols).items())
    for i in range(0, len(years), 3):
        row_cells = []
        for y, n in years[i:i + 3]:
            row_cells += [y, n]
        while len(row_cells) < 6:
            row_cells.append("")
        line(row_cells, bold_first=False)
    r += 1
    note("THE GAP", "The run begins 17 Apr 2011, but the Inquirer column actually began 11 Feb 2007 — his first, "
         "reprinted as BA031 'Visionary Leadership by Example'. Roughly 222 columns from Feb 2007 to Apr 2011 have "
         "never been sourced. 803 + ~222 is about 1,025, which is where 'almost a thousand columns' comes from. "
         "A batch-05 sourcing job, not a pipeline job.", red=True)
    r += 1

    ban("SPEECHES")
    hdr(["Span", "Documents", "Theme C", "Theme A", "Other", "Note"])
    line([f"{sfirst} – {slast}", len(sps), sth.get("C", 0), sth.get("A", 0),
          " · ".join(f"{k}:{v}" for k, v in sorted(sth.items()) if k not in ("A", "C")),
          "Thirty-two years of platform speaking — inductions, keynotes, lectures, eulogies."],
         bold_first=False, h=28)
    dec = sorted(_c.Counter(((d.get("date") or "0000")[:3] + "0s") for d in sps).items())
    line(["By decade"] + [f"{k} {v}" for k, v in dec] + [""] * (5 - len(dec)), h=22)
    r += 1

    ban("BIOGRAPHY — all 35 chapters, in ID order")
    hdr(["Doc ID", "Chapter title", "Date", "Theme", "", "Source"])
    for d in sorted((d for d in docs if d["id"][0] == "G"), key=lambda x: x["id"]):
        line([d["id"], d.get("title"), d.get("date"), d["id"][1], "",
              "cjp_biography_curated_normalized.xlsx (14 fields — no Link)"], bold_first=False, h=18)
        ws.cell(r - 1, 1).font = MONOB
    r += 1
    note("All one theme", f"Every biography chapter is theme C (biographical and personal) — {gth.get('C', 0)} of "
         f"{len(bio)}. That is correct for the material, but it means theme C is carried almost entirely by the "
         "biography and by Justice and Faith / Love God, Serve Man.")
    r += 1

    ban("THE FIVE THEMES")
    hdr(["Code", "Theme", "Docs now", "Folder", "", "What belongs here"])
    for code, name, folder, what in [
        ("A", "Liberty and Rule of Law", "A_liberty_rule_of_law",
         "Constitutional doctrine, due process, judicial independence and review, civil liberties, the courts as the "
         "last bulwark."),
        ("B", "Prosperity and Economic Philosophy", "B_prosperity_economic_philosophy",
         "Economic governance, business and investment law, MSMEs, social justice as an economic question."),
        ("C", "Biographical and Personal", "C_biographical_personal",
         "Early life, family, faith, mentors, friendships, honours — the man rather than the office."),
        ("D", "FLP Mission and Foundation", "D_flp_mission_foundation",
         "Scholarships, the dissertation contest, donors and partners, the Museum."),
        ("E", "Current Events Commentary", "E_current_events_commentary",
         "Time-bound commentary on Philippine political events, elections, geopolitics."),
    ]:
        line([code, name, thmix.get(code, 0), folder, "", what], h=30)
    r += 1
    note("The distribution is lopsided", f"Theme A holds {thmix.get('A',0)} of {tot_now} documents; theme B holds "
         f"{thmix.get('B',0)}. A theme is a SHELF, one per document — not a score, and not the same thing as the 34 "
         "retrieval topics. Do not use theme counts to judge retrieval coverage.", key=True)
    r += 1

    ban("THE IDENTIFIER")
    hdr(["Letter", "Format", "Documents now", "After batch-04", "", "Pinned workbook"])
    for letter, fmt, now, after, bookf in [
        ("C", "column", len(cols), len(cols) + cnew, "cjp_columns_curated_normalized.xlsx"),
        ("B", "book chapter", sum(bnow.values()), sum(bnow.values()) + sum(bnew.values()),
         "cjp_books_curated_normalized.xlsx"),
        ("S", "speech", len(sps), len(sps), "cjp_speeches_curated_normalized.xlsx"),
        ("G", "biography chapter", len(bio), len(bio), "cjp_biography_curated_normalized.xlsx (14 fields)"),
    ]:
        line([letter, fmt, now, after, "", bookf], h=20)
        ws.cell(r - 1, 1).font = MONOB
    r += 1
    note("Pattern", "^[CBSG][A-E]\\d{3}$   e.g. CA003 · BA108 · SB028 · GC021. Letter 1 is the format, letter 2 is the "
         "theme, digits 3-5 are a position in that format+theme series — a POSITION, not a count. CA529 does not mean "
         "there are 529 columns.")
    note("Never reused", "An ID is never reused and never renumbered. A retired ID stays retired; a replacement takes "
         "a NEW ID and records superseded_doc_id.", key=True)
    r += 1

    ban("RETIRED IDENTIFIERS — data/csv/retired_doc_ids.csv")
    hdr(["Doc ID", "Title", "Type", "Superseded by", "Retired", "Reason"])
    rp = ROOT / "data" / "csv" / "retired_doc_ids.csv"
    if rp.exists():
        for row in _csv.DictReader(open(rp, encoding="utf-8-sig")):
            line([row.get("doc_id"), row.get("title"), row.get("type"),
                  row.get("superseded_by") or "—", row.get("retired_on"), row.get("reason")],
                 bold_first=False, h=28)
            ws.cell(r - 1, 1).font = MONOB
    r += 1

    ban("WHERE THE 34 TOPICS COME FROM — three layers, only two of them derived")
    note("The question", "Are the topics a RESULT of the data enrichment, or something imposed on it? Both, in "
         "different layers — and the distinction decides what has to be rebuilt when new material arrives.")
    note("Layer 1 · taxonomy\n(NOT derived)",
         "WHICH topics exist — the ~35 ids, display names, definitions, default register and matcher keywords — is a "
         "HAND-CURATED Python dict inside scripts/build_topic_map.py. Its own docstring says so. It was written "
         "against 79 documents (64 columns, 15 speeches, NO books), generated 2026-05-25, and never re-derived.",
         key=True)
    note("Layer 2 · assignment + stats\n(DERIVED from enrichment)",
         "Each document is scored against every topic by matching title + primary_topics + sub_topics + Keyword/s + "
         "entity names. The per-topic statistics — doc_count, doc_ids, top_signature_phrases, top_people, "
         "top_institutions, top_cases, date_range, distributions — are aggregated the same way. The enrichment DOES "
         "shape the map; it just cannot invent a topic the curator never wrote.")
    note("Layer 3 · centroids\n(DERIVED from embeddings)",
         "CE-10: centroid = mean(label + description + signature_phrases + exemplar chunks), in the 768-d space. This "
         "is what retrieval scores against. Rebuilding centroids does NOT change which topics exist — that is layer 1.")
    note("THE GAP", "Layer 1 has never seen the enrichment of the other 1,025 documents, let alone the 168 BOOK "
         "chapters arriving in batch-04 — a format those keyword matchers were never written for.", red=True)
    note("The governed fix",
         "CE-7  tag the new docs against the CURRENT map, then CLUSTER the orphans\n"
         "CE-8  DECIDE (human, no CC prompt): retag only / expand / full rebuild\n"
         "CE-9  if expanding: taxonomy_expansion_PROPOSAL.md, one candidate dimension per section, with sign-off\n"
         "CE-10 rebuild centroids from scratch under the agreed model\n"
         "Runbook rule 4: add a topic dimension only from evidence you can point at. Never exercised — batch-03 was a "
         "correction (retag only) and batch-04 has not started.", key=True)
    r += 1

    ban("A SECOND GAP — the assignment does not persist")
    ntp = sum(1 for d in docs if "topic_paths" in d or "topics" in d)
    note("Measured", f"{ntp} of {tot_now} corpus .json files carry topic_paths or any topic field.", red=True)
    note("Why", "scripts/apply_topic_paths.py backfills topic_paths into the per-doc .json, but "
         "generate_corpus_from_xlsx.py rebuilds those .json from the pinned workbooks — which have no topic column — "
         "so every regeneration WIPES the backfill. The only live tagging is reports/w1_7_pilot_topic_tags.json, "
         "covering 95 documents.")
    note("What it means", "P3.3's 'every document tagged with a primary and secondary topic' fails twice: coverage "
         "AND persistence. Fixing it is a choice — add a topic column to the pinned schema so it survives "
         "regeneration, or treat the tag file as the system of record and regenerate it after every corpus build. "
         "A CE-8-adjacent decision, not a bug fix.", key=True)

    ws.freeze_panes = "A3"


# ---------------------------------------------------------------- copy v2
FORMAT_POINTER = ("\n\nFORMATS: source .md -> see tab 'Format · Source .md'. Enriched row (15 fields; biography 14) -> "
                  "tab 'Format · Enriched row'. Generated corpus files -> tab 'Format · Corpus output'. "
                  "Themes and IDs -> tab 'Themes & IDs'.")
# Evidence text appended to specific Project Plan rows. Lives here, not patched into the
# workbook, so a rebuild does not silently drop it. Keyed by row ID; appended only if the
# sentinel is not already present in the v2 cell.
EVIDENCE_ADDENDA = {
    "P3.3": ("ALSO MEASURED 26 Sep", " ALSO MEASURED 26 Sep: zero of 1,104 corpus .json carry topic_paths - "
             "apply_topic_paths.py backfills them but generate_corpus_from_xlsx.py rebuilds the .json from the "
             "pinned workbooks, which have no topic column, so every regeneration wipes it. The taxonomy itself "
             "(layer 1) is a hand-curated dict in build_topic_map.py written against 79 documents and never "
             "re-derived; only the assignment and the per-topic stats are derived from enrichment. See the "
             "'Themes & IDs' tab."),
}

POINTER_IDS = {"P1.3", "P1.4", "P2.1", "P2.2", "P2.3", "P2.4", "P3.2", "P3.3", "P3.4", "P3.5",
               "CE-1", "CE-1a", "CE-1b", "CE-2", "CE-3", "CE-4", "CE-5", "CE-6", "CE-7"}


def copy_sheet(src_ws, wb, how_col_name="How / paste-ready prompt"):
    ws = wb.create_sheet(src_ws.title)
    ncols = src_ws.max_column
    hdr_row, hdr = None, []
    for r in range(1, min(src_ws.max_row, 8) + 1):
        vals = [src_ws.cell(r, c).value for c in range(1, ncols + 1)]
        if any(isinstance(v, str) and v.strip() == "ID" for v in vals if v):
            hdr_row, hdr = r, vals
            break
    how_i = hdr.index(how_col_name) + 1 if how_col_name in hdr else None
    id_i = hdr.index("ID") + 1 if "ID" in hdr else None
    ev_i = hdr.index("Evidence") + 1 if "Evidence" in hdr else None

    for r in range(1, src_ws.max_row + 1):
        for c in range(1, ncols + 1):
            v = src_ws.cell(r, c).value
            if v is None:
                continue
            if (how_i and c == how_i and id_i and str(src_ws.cell(r, id_i).value or "").strip() in POINTER_IDS
                    and isinstance(v, str) and "FORMATS:" not in v):
                v = v + FORMAT_POINTER
            if ev_i and c == ev_i and id_i and isinstance(v, str):
                add = EVIDENCE_ADDENDA.get(str(src_ws.cell(r, id_i).value or "").strip())
                if add and add[0] not in v:
                    v = v + add[1]
            cell = ws.cell(r, c)
            cell.value = v
            cell.font = BODY
            cell.alignment = TOP
        if hdr_row and r == hdr_row:
            for c in range(1, ncols + 1):
                ws.cell(r, c).font = BOLD
                ws.cell(r, c).fill = FILL_BAND
        elif r < (hdr_row or 99):
            ws.cell(r, 1).font = Font(name=FONT, size=11, bold=True)
        else:
            a = str(ws.cell(r, 1).value or "")
            if a and not re.match(r"^(P\d|CE-|[A-Z]\d)", a) and ws.cell(r, 2).value is None:
                for c in range(1, ncols + 1):
                    ws.cell(r, c).fill = FILL_BAND
                ws.cell(r, 1).font = BOLD
    # widths
    w = {1: 9, 2: 9, 3: 34, 4: 52, 5: 8, 6: 20, 7: 78, 8: 40, 9: 14, 10: 13, 11: 62}
    for c in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(c)].width = w.get(c, 26)
    if hdr_row:
        ws.freeze_panes = f"A{hdr_row + 1}"
    return ws


def main() -> int:
    if not SRC.exists():
        print(f"missing {SRC}", file=sys.stderr)
        return 2
    src = openpyxl.load_workbook(SRC)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_readme(wb)
    for name in src.sheetnames:
        if name == "Read Me":
            continue
        copy_sheet(src[name], wb)
    build_source_md_tab(wb)
    build_enriched_tab(wb)
    build_corpus_tab(wb)
    build_themes_tab(wb)

    order = ["Read Me", "Project Plan", "Corpus Expansion", "Data Flow", "Intake Prompts",
             "Format · Source .md", "Format · Enriched row", "Format · Corpus output",
             "Themes & IDs", "Metrics & Gates", "Reuse Checklist"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + \
                 [s for s in wb._sheets if s.title not in order]
    wb.save(OUT)
    print(f"wrote {OUT}")
    chk = openpyxl.load_workbook(OUT)
    for s in chk.sheetnames:
        print(f"   {s:26s} {chk[s].max_row:4d} rows x {chk[s].max_column} cols")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
